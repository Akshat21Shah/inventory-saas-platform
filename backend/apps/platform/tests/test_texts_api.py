"""Translations for the super admin (ADR-060 items 11, 14): each language's review progress, the
translation sheet to download, and the better words staff and shops suggest from any screen.
Only the super admin sees the list and the sheet; anyone signed in may suggest."""

from io import BytesIO
from typing import Any

import pytest
from django.core.cache import cache
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.tests.factories import make_staff_in, make_super_admin
from apps.accounts.tokens import issue_tokens
from apps.audit.models import AuditLog
from apps.orders.tests.helpers import client_for, make_shop, shop_client
from apps.platform import text_suggestions
from apps.platform.models import TextSuggestion
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"
SUGGEST = f"{API}/texts/suggestions/"
LIST = f"{API}/platform/texts/suggestions/"


@pytest.fixture(autouse=True)
def _clean():
    cache.clear()
    yield
    cache.clear()


@pytest.fixture
def admin() -> APIClient:
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {issue_tokens(make_super_admin(), None).access}")
    client.defaults["HTTP_X_FORWARDED_HOST"] = "admin.localhost"
    return client


def _word(**extra: Any) -> dict[str, str]:
    return {
        "language": "hi",
        "screen": "/shop/orders",
        "current_text": "बकाया",
        "suggestion": "उधार",
        **extra,
    }


@covers("text-suggestions")
def test_staff_shops_and_the_super_admin_may_suggest_a_better_word(tenant_a, admin):
    staff = client_for(tenant_a, make_staff_in(tenant_a, "SALES"))
    assert staff.post(SUGGEST, _word(), format="json").status_code == 201
    shop = shop_client(tenant_a, make_shop(tenant_a, "9876500161"))
    assert shop.post(SUGGEST, _word(language="mr"), format="json").status_code == 201
    assert admin.post(SUGGEST, _word(), format="json").status_code == 201
    rows = list(TextSuggestion.objects.order_by("created_at"))
    assert [r.tenant_id for r in rows] == [tenant_a.pk, tenant_a.pk, None]
    assert APIClient().post(SUGGEST, _word(), format="json").status_code == 401
    refused = staff.post(SUGGEST, _word(language="xx", suggestion=""), format="json")
    assert refused.status_code == 400


def test_suggestions_are_limited_per_hour(tenant_a):
    staff = client_for(tenant_a, make_staff_in(tenant_a, "SALES"))
    for _ in range(text_suggestions.PER_HOUR):
        assert staff.post(SUGGEST, _word(), format="json").status_code == 201
    over = staff.post(SUGGEST, _word(), format="json")
    assert (over.status_code, over.json()["error"]["code"]) == (429, "RATE_LIMITED")


@covers("platform-text-suggestions", "platform-text-suggestion")
def test_the_super_admin_lists_and_settles_them_and_nobody_else_can(tenant_a, tenant_b, admin):
    owner_a = client_for(tenant_a, make_staff_in(tenant_a, "OWNER"))
    owner_b = client_for(tenant_b, make_staff_in(tenant_b, "OWNER"))
    owner_a.post(SUGGEST, _word(), format="json")
    owner_b.post(SUGGEST, _word(language="mr", suggestion="थकबाकी"), format="json")
    listed = admin.get(LIST).json()["results"]
    assert {(r["tenant_name"], r["language"]) for r in listed} == {
        (tenant_a.name, "hi"),
        (tenant_b.name, "mr"),
    }
    assert [r["suggestion"] for r in admin.get(LIST, {"language": "mr"}).json()["results"]] == [
        "थकबाकी"
    ]
    # A distributor sees no one's list, not even its own people's.
    assert owner_a.get(LIST).status_code == 403
    pk = listed[0]["id"]
    assert owner_a.patch(f"{LIST}{pk}/", {"status": "DONE"}, format="json").status_code == 403
    done = admin.patch(f"{LIST}{pk}/", {"status": "DONE"}, format="json")
    assert done.status_code == 200 and done.json()["status"] == "DONE"
    assert done.json()["resolved_by"] is not None
    assert [r["id"] for r in admin.get(LIST, {"status": "NEW"}).json()["results"]] == [
        listed[1]["id"]
    ]
    assert AuditLog.objects.filter(action="platform.text_suggestion_resolved").exists()


@covers("platform-texts-progress", "platform-texts-sheet")
def test_the_super_admin_sees_progress_and_downloads_the_sheet(tenant_a, admin):
    client_for(tenant_a, make_staff_in(tenant_a, "OWNER")).post(SUGGEST, _word(), format="json")
    progress = {row["code"]: row for row in admin.get(f"{API}/platform/texts/progress/").json()}
    assert set(progress) == {"hi", "mr"}
    hindi = progress["hi"]
    assert hindi["enabled"] is False  # off for everyone until the review (item 42)
    assert hindi["total"] > 5000 and hindi["translated"] == hindi["total"]
    assert (hindi["reviewed"], hindi["new_suggestions"], progress["mr"]["new_suggestions"]) == (
        0,
        1,
        0,
    )
    sheet = admin.get(f"{API}/platform/texts/sheet/")
    assert sheet.status_code == 200
    assert sheet["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in sheet["Content-Disposition"]
    book = load_workbook(BytesIO(sheet.content), read_only=True)
    assert book.sheetnames == ["Texts", "How to use"]
    assert AuditLog.objects.filter(action="platform.texts_exported").exists()
    owner = client_for(tenant_a, make_staff_in(tenant_a, "OWNER"))
    assert owner.get(f"{API}/platform/texts/sheet/").status_code == 403
    assert owner.get(f"{API}/platform/texts/progress/").status_code == 403
