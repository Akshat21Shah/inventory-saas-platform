"""Per-shop pricing tools (ADR-037): discount grid, copy pricing, bulk % price-list change, shop
pricing report and free products. Behaviour, audit, stale previews and tenant isolation."""

import io
from datetime import timedelta
from decimal import Decimal as D

import pytest
from django.core.cache import cache
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.tests.factories import make_staff_in
from apps.accounts.tokens import issue_tokens
from apps.audit.models import AuditLog
from apps.catalog.models import Brand, Category, Product, ProductTaxRate, Unit
from apps.pricing.models import DiscountRule, DiscountSlab, PriceList, PriceListItem, RetailerPrice
from apps.retailers.services import create_retailer
from common.dates import today_ist
from common.tenancy import tenant_context
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"


@pytest.fixture(autouse=True)
def _clean():
    cache.clear()
    yield
    cache.clear()


def _client(tenant, role="OWNER"):
    client = APIClient()
    user = make_staff_in(tenant, role)
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {issue_tokens(user, tenant.pk).access}")
    client.defaults["HTTP_X_FORWARDED_HOST"] = f"{tenant.slug}.localhost"
    return client


@pytest.fixture
def w(tenant_a):
    with tenant_context(tenant_a.pk):
        unit = Unit.objects.get(code="PCS")
        food = Category.objects.create(name="Food", slug="food")
        lays = Brand.objects.create(name="Lays")

        def product(code, price, **extra):
            p = Product.objects.create(
                code=code,
                name=f"Item {code}",
                unit=unit,
                hsn_code="1905",
                base_price=D(price),
                **extra,
            )
            ProductTaxRate.objects.create(product=p, gst_rate=D("5"), effective_from=today_ist())
            return p

        a = product("A", "10.00", brand=lays, category=food)
        b = product("B", "20.00", brand=lays)
        c = product("C", "33.33", category=food)
        gold = PriceList.objects.create(name="Gold")
        PriceListItem.objects.create(price_list=gold, product=a, price=D("9.00"))
        silver = PriceList.objects.create(name="Silver")
        src = create_retailer(
            shop_name="Source",
            phone="9876500001",
            send_welcome=False,
            extra={"price_list_id": gold.pk},
        )
        dst = create_retailer(
            shop_name="Target",
            phone="9876500002",
            send_welcome=False,
            extra={"price_list_id": silver.pk},
        )
    return {
        "a": a,
        "b": b,
        "c": c,
        "food": food,
        "lays": lays,
        "gold": gold,
        "silver": silver,
        "src": src,
        "dst": dst,
    }


# --- Discount grid -----------------------------------------------------------------------------


@covers("retailer-discount-grid", "retailer-discount-grid-preview")
def test_discount_grid(tenant_a, tenant_b, w):
    owner = _client(tenant_a)
    shop = w["dst"]
    with tenant_context(tenant_a.pk):
        dated = DiscountRule.objects.create(
            name="Dated",
            discount_type="PERCENT",
            value=D("3"),
            scope_type="PRODUCT",
            product=w["a"],
            audience_type="RETAILER",
            retailer=shop,
            valid_to=today_ist() + timedelta(days=5),
        )
    url = f"{API}/retailers/{shop.pk}/discount-grid/"
    rows = {r["product"]["code"]: r for r in owner.get(url).json()["results"]}
    assert set(rows) == {"A", "B", "C"}
    assert rows["A"]["simple"] is None and [o["name"] for o in rows["A"]["others"]] == ["Dated"]
    assert rows["B"]["price"]["net_unit_price"] == "20.00"
    assert [
        r["product"]["code"] for r in owner.get(url, {"brand": str(w["lays"].pk)}).json()["results"]
    ] == ["A", "B"]

    # Live preview: nothing is saved.
    preview = owner.post(
        f"{url}preview/",
        {
            "items": [
                {"product": str(w["b"].pk), "discount_type": "PERCENT", "value": "10"},
            ]
        },
        format="json",
    ).json()
    assert preview[0]["price"]["net_unit_price"] == "18.00"
    assert preview[0]["price"]["discount_percent"] == "10.00"
    with tenant_context(tenant_a.pk):
        assert DiscountRule.objects.count() == 1

    saved = owner.put(
        url,
        {
            "items": [
                {"product": str(w["b"].pk), "discount_type": "PERCENT", "value": "10"},
                {"product": str(w["c"].pk), "discount_type": "FLAT_PER_UNIT", "value": "1.50"},
            ]
        },
        format="json",
    ).json()
    assert saved == {"changed": 2, "warnings": []}
    changed = owner.put(
        url,
        {
            "items": [
                {"product": str(w["b"].pk), "discount_type": "PERCENT", "value": "12.5"},
                {"product": str(w["c"].pk), "discount_type": "FLAT_PER_UNIT", "value": None},
            ]
        },
        format="json",
    ).json()
    assert changed["changed"] == 2
    entry = AuditLog.objects.filter(action="pricing.shop_discounts_changed").last()
    assert entry is not None and entry.changes == {"B": ["10%", "12.5%"], "C": ["₹1.50 each", None]}
    with tenant_context(tenant_a.pk):
        simple = DiscountRule.objects.get(product=w["b"])
        assert (simple.name, simple.value, simple.retailer_id) == (
            f"{shop.code} · B",
            D("12.50"),
            shop.pk,
        )
        assert not DiscountRule.objects.filter(product=w["c"]).exists()
        assert DiscountRule.objects.filter(pk=dated.pk).exists()  # the dated rule is untouched
    rows = {r["product"]["code"]: r for r in owner.get(url).json()["results"]}
    assert rows["B"]["simple"]["value"] == "12.50"
    assert rows["B"]["price"]["net_unit_price"] == "17.50"

    free = owner.put(
        url,
        {
            "items": [
                {"product": str(w["a"].pk), "discount_type": "PERCENT", "value": "100"},
            ]
        },
        format="json",
    ).json()
    assert free["warnings"][0]["message"].startswith("These discounts make 1 product free")
    bad = owner.put(
        url,
        {
            "items": [
                {"product": str(w["a"].pk), "discount_type": "PERCENT", "value": "120"},
            ]
        },
        format="json",
    )
    assert "items.0.value" in bad.json()["error"]["details"]["fields"]

    other = _client(tenant_b)
    assert other.get(url).status_code == 404
    assert (
        other.put(
            url,
            {"items": [{"product": str(w["a"].pk), "discount_type": "PERCENT", "value": "5"}]},
            format="json",
        ).status_code
        == 404
    )
    assert (
        other.post(
            f"{url}preview/",
            {"items": [{"product": str(w["a"].pk), "discount_type": "PERCENT", "value": "5"}]},
            format="json",
        ).status_code
        == 404
    )
    sales = _client(tenant_a, "SALES")  # pricing.view: may look, may not save
    assert sales.get(url).status_code == 200
    assert sales.put(url, {"items": []}, format="json").status_code == 403


# --- Copy pricing ------------------------------------------------------------------------------


def _pricing_of(tenant, shop):
    with tenant_context(tenant.pk):
        shop.refresh_from_db()
        return (
            shop.price_list.name if shop.price_list else None,
            sorted((p.product.code, p.price) for p in RetailerPrice.objects.filter(retailer=shop)),
            sorted(r.name for r in DiscountRule.objects.filter(retailer=shop)),
        )


@pytest.fixture
def two_shops(tenant_a, w):
    src, dst = w["src"], w["dst"]
    with tenant_context(tenant_a.pk):
        RetailerPrice.objects.create(retailer=src, product=w["a"], price=D("8.00"))
        RetailerPrice.objects.create(retailer=src, product=w["b"], price=D("18.00"))
        RetailerPrice.objects.create(retailer=dst, product=w["b"], price=D("19.00"))
        RetailerPrice.objects.create(retailer=dst, product=w["c"], price=D("30.00"))
        slabbed = DiscountRule.objects.create(
            name=f"{src.code} bulk",
            discount_type="PERCENT",
            value=D("0"),
            scope_type="CATEGORY",
            category=w["food"],
            audience_type="RETAILER",
            retailer=src,
        )
        DiscountSlab.objects.create(rule=slabbed, min_qty=D("12"), value=D("4"))
        DiscountRule.objects.create(
            name=f"{src.code} · A",
            discount_type="PERCENT",
            value=D("5"),
            scope_type="PRODUCT",
            product=w["a"],
            audience_type="RETAILER",
            retailer=src,
        )
        DiscountRule.objects.create(
            name="Target's own",
            discount_type="PERCENT",
            value=D("2"),
            scope_type="ALL",
            audience_type="RETAILER",
            retailer=dst,
        )
    return src, dst


@covers("retailer-copy-pricing-preview", "retailer-copy-pricing")
@pytest.mark.parametrize("mode", ["REPLACE", "ADD"])
def test_copy_pricing(tenant_a, tenant_b, w, two_shops, mode):
    src, dst = two_shops
    owner = _client(tenant_a)
    base = f"{API}/retailers/{dst.pk}/copy-pricing/"
    body = {"copy_from": str(src.pk), "mode": mode}
    plan = owner.post(f"{base}preview/", body, format="json").json()
    assert (plan["price_list_from"], plan["price_list_to"]) == ("Silver", "Gold")
    assert [p["code"] for p in plan["prices_add"]] == ["A"]
    assert [(p["code"], p["old"], p["new"]) for p in plan["prices_update"]] == [
        ("B", "19.00", "18.00")
    ]
    if mode == "REPLACE":
        assert [p["code"] for p in plan["prices_remove"]] == ["C"]
        assert plan["rules_remove"] == ["Target's own"]
        assert plan["changes"] == 1 + 1 + 1 + 1 + 2 + 1
    else:
        assert plan["prices_remove"] == [] and plan["rules_remove"] == []

    stale = owner.post(base, {**body, "expected_changes": plan["changes"] + 1}, format="json")
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "PREVIEW_OUT_OF_DATE"
    done = owner.post(base, {**body, "expected_changes": plan["changes"]}, format="json")
    assert done.status_code == 200, done.json()

    price_list, prices, rules = _pricing_of(tenant_a, dst)
    assert price_list == "Gold"
    copied = [f"{dst.code} bulk", f"{dst.code} · A"]
    if mode == "REPLACE":
        assert prices == [("A", D("8.00")), ("B", D("18.00"))]
        assert rules == sorted(copied)
    else:
        assert prices == [("A", D("8.00")), ("B", D("18.00")), ("C", D("30.00"))]
        assert rules == sorted([*copied, "Target's own"])
    with tenant_context(tenant_a.pk):
        copy = DiscountRule.objects.get(name=f"{dst.code} bulk")
        assert [(s.min_qty, s.value) for s in copy.slabs.all()] == [(D("12.000"), D("4.00"))]
    entry = AuditLog.objects.get(action="pricing.pricing_copied")
    assert entry.metadata["mode"] == mode and entry.metadata["from"].startswith(src.code)
    assert _pricing_of(tenant_a, src)[1] == [("A", D("8.00")), ("B", D("18.00"))]  # untouched

    no_mode = owner.post(f"{base}preview/", {"copy_from": str(src.pk)}, format="json")
    assert no_mode.json()["error"]["details"]["fields"]["mode"] == ["Choose “Replace” or “Add”."]
    same = owner.post(f"{base}preview/", {"copy_from": str(dst.pk), "mode": mode}, format="json")
    assert "copy_from" in same.json()["error"]["details"]["fields"]
    other = _client(tenant_b)
    assert other.post(f"{base}preview/", body, format="json").status_code == 404
    assert other.post(base, {**body, "expected_changes": 0}, format="json").status_code == 404


def test_copy_in_add_mode_replaces_a_shop_and_product_simple_rule(tenant_a, w, two_shops):
    src, dst = two_shops
    with tenant_context(tenant_a.pk):
        DiscountRule.objects.create(
            name=f"{dst.code} · A",
            discount_type="PERCENT",
            value=D("1"),
            scope_type="PRODUCT",
            product=w["a"],
            audience_type="RETAILER",
            retailer=dst,
        )
    owner = _client(tenant_a)
    base = f"{API}/retailers/{dst.pk}/copy-pricing/"
    plan = owner.post(
        f"{base}preview/", {"copy_from": str(src.pk), "mode": "ADD"}, format="json"
    ).json()
    assert plan["rules_replace"] == [f"{dst.code} · A"]
    owner.post(
        base,
        {"copy_from": str(src.pk), "mode": "ADD", "expected_changes": plan["changes"]},
        format="json",
    )
    with tenant_context(tenant_a.pk):
        [rule] = DiscountRule.objects.filter(retailer=dst, product=w["a"])
        assert rule.value == D("5.00")  # one simple rule per shop and product


# --- Bulk % change to a price list -------------------------------------------------------------


@covers("price-list-adjust-preview", "price-list-adjust")
def test_bulk_percentage_change(tenant_a, tenant_b, w):
    owner = _client(tenant_a)
    gold = w["gold"]
    with tenant_context(tenant_a.pk):
        PriceListItem.objects.create(price_list=gold, product=w["b"], price=D("19.99"))
    base = f"{API}/price-lists/{gold.pk}/adjust/"
    body = {"percent": "5", "brand": str(w["lays"].pk)}
    preview = owner.post(f"{base}preview/", body, format="json").json()
    assert [(r["code"], r["old"], r["new"]) for r in preview["rows"]] == [
        ("A", "9.00", "9.45"),
        ("B", "19.99", "20.99"),
    ]  # 20.9895 → 20.99
    rupees = owner.post(f"{base}preview/", {**body, "rounding": "RUPEE"}, format="json").json()
    # 9.45 → 9 (no change, so not listed); 20.99 → 21
    assert [(r["code"], r["new"]) for r in rupees["rows"]] == [("B", "21.00")]
    assert rupees["count"] == 1
    more = owner.post(
        f"{base}preview/",
        {"percent": "-10", "category": str(w["food"].pk), "include_missing": True},
        format="json",
    ).json()
    assert [(r["code"], r["old"], r["new"]) for r in more["rows"]] == [
        ("A", "9.00", "8.10"),
        ("C", None, "30.00"),
    ]  # C starts from the standard 33.33
    assert more["added"] == 1

    stale = owner.post(base, {**body, "expected_count": 5}, format="json")
    assert stale.json()["error"]["code"] == "PREVIEW_OUT_OF_DATE"
    done = owner.post(base, {**body, "expected_count": 2}, format="json").json()
    assert done == {"changed": 2, "warnings": []}
    with tenant_context(tenant_a.pk):
        prices = dict(
            PriceListItem.objects.filter(price_list=gold).values_list("product__code", "price")
        )
    assert prices == {"A": D("9.45"), "B": D("20.99")}
    assert AuditLog.objects.get(action="pricing.price_list_adjusted").metadata["percent"] == "5"
    assert AuditLog.objects.filter(action="pricing.price_list_prices_changed").exists()

    both = owner.post(f"{base}preview/", {"percent": "5"}, format="json")
    assert both.json()["error"]["details"]["fields"]["scope"] == ["Choose a category or a brand."]
    other = _client(tenant_b)
    assert other.post(f"{base}preview/", body, format="json").status_code == 404
    assert other.post(base, {**body, "expected_count": 2}, format="json").status_code == 404


# --- Report and free products ------------------------------------------------------------------


@covers("pricing-shop-report", "pricing-shop-report-export", "retailer-free-products")
def test_shop_pricing_report(tenant_a, tenant_b, w, two_shops):
    src, dst = two_shops
    with tenant_context(tenant_a.pk):
        plain = create_retailer(shop_name="Plain", phone="9876500003", send_welcome=False)
        RetailerPrice.objects.create(retailer=dst, product=w["a"], price=D("0"))  # free
        DiscountRule.objects.create(  # makes B free for the source shop (18 - 100%)
            name="B free",
            discount_type="PERCENT",
            value=D("100"),
            scope_type="PRODUCT",
            product=w["b"],
            audience_type="RETAILER",
            retailer=src,
        )
    owner = _client(tenant_a)
    report = owner.get(f"{API}/pricing/shop-report/").json()["results"]
    rows = {r["shop_name"]: r for r in report}
    assert set(rows) == {"Source", "Target"}  # only shops with their own pricing
    assert (rows["Source"]["special_price_count"], rows["Source"]["shop_rule_count"]) == (2, 3)
    assert rows["Source"]["free_product_count"] == 1
    assert rows["Target"]["free_product_count"] == 1
    everyone = owner.get(f"{API}/pricing/shop-report/", {"show": "ALL"}).json()["results"]
    assert "Plain" in {r["shop_name"] for r in everyone}
    free = owner.get(f"{API}/retailers/{dst.pk}/free-products/").json()
    assert [p["code"] for p in free] == ["A"]
    assert owner.get(f"{API}/retailers/{plain.pk}/free-products/").json() == []

    book = load_workbook(io.BytesIO(owner.get(f"{API}/pricing/shop-report/export/").content))
    sheet = list(book.worksheets[0].iter_rows(values_only=True))
    assert sheet[0][0] == "Shop code" and len(sheet) == 3

    other = _client(tenant_b)
    assert other.get(f"{API}/pricing/shop-report/", {"show": "ALL"}).json()["results"] == []
    assert other.get(f"{API}/retailers/{dst.pk}/free-products/").status_code == 404
    other_book = load_workbook(io.BytesIO(other.get(f"{API}/pricing/shop-report/export/").content))
    assert len(list(other_book.worksheets[0].iter_rows())) == 1
    warehouse = _client(tenant_a, "WAREHOUSE")  # no pricing.view
    assert warehouse.get(f"{API}/pricing/shop-report/").status_code == 403


@pytest.mark.parametrize(("combination", "free"), [("BEST", False), ("ADD", True)])
def test_free_products_follow_the_combination_setting(tenant_a, w, combination, free):
    from apps.platform.models import TenantSetting
    from apps.pricing.tools import free_products

    with tenant_context(tenant_a.pk):
        TenantSetting.objects.update_or_create(
            key="pricing.discount_combination", defaults={"value": combination}
        )
        for value in ("60", "40"):
            DiscountRule.objects.create(
                name=value,
                discount_type="PERCENT",
                value=D(value),
                scope_type="PRODUCT",
                product=w["c"],
                audience_type="ALL",
            )
    cache.clear()
    with tenant_context(tenant_a.pk):
        found = [p.code for p in free_products(w["dst"])]
    assert found == (["C"] if free else [])
