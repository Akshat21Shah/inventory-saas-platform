"""Pricing imports and exports (ADR-037 item 6): special prices, price-list prices and discount
rules through the import framework (explicit mode, change preview, plain row/column errors),
round trips, permissions and tenant isolation."""

import io
from datetime import date
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.catalog.models import Brand, Category, Product, ProductTaxRate, Unit
from apps.dataio.tests.helpers import _client, commit, messages, upload, xlsx
from apps.pricing.models import DiscountRule, PriceList, PriceListItem, RetailerPrice
from apps.retailers.services import create_retailer
from common.dates import today_ist
from common.tenancy import tenant_context
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"


@pytest.fixture
def world(tenant_a):
    with tenant_context(tenant_a.pk):
        unit = Unit.objects.get(code="PCS")
        food = Category.objects.create(name="Food", slug="food")
        biscuits = Category.objects.create(name="Biscuits", slug="biscuits", parent=food, level=2)
        parle = Brand.objects.create(name="Parle")
        products = {}
        for code, price in (("PG-100", "10.00"), ("MG-70", "12.50")):
            products[code] = Product.objects.create(
                code=code,
                name=f"Item {code}",
                unit=unit,
                hsn_code="1905",
                base_price=D(price),
                category=biscuits,
                brand=parle,
            )
            ProductTaxRate.objects.create(
                product=products[code], gst_rate=D("5"), effective_from=today_ist()
            )
        gold = PriceList.objects.create(name="Gold")
        shop = create_retailer(shop_name="Ganesh", phone="9876500001", send_welcome=False)
        other = create_retailer(shop_name="Laxmi", phone="9876500002", send_welcome=False)
    return {
        "products": products,
        "gold": gold,
        "shop": shop,
        "other": other,
        "food": food,
        "biscuits": biscuits,
        "parle": parle,
    }


def _export(client, path):
    response = client.get(f"{API}/{path}", {"file_type": "xlsx"})
    assert response.status_code == 200, response.content[:300]
    return [
        list(r)
        for r in load_workbook(io.BytesIO(response.content))
        .worksheets[0]
        .iter_rows(values_only=True)
    ]


def _file(rows):
    return xlsx([[("" if c is None else c) for c in row] for row in rows])


# --- Special prices ------------------------------------------------------------------------------


@covers("special-prices-export")
def test_special_prices_import_and_export(tenant_a, tenant_b, world, owner, run):
    shop = world["shop"]
    header = ["Shop", "Product code", "Special price", "Note"]
    job = upload(
        owner,
        run,
        _file(
            [
                header,
                ["98765 00001", "PG-100", "₹8.75", "Deal"],  # by mobile, as typed
                [shop.code, "MG-70", "0", ""],  # by code; a ₹0 price warns
                ["9999999999", "PG-100", "5", ""],
                [shop.code, "NOPE", "5", ""],
                [shop.code, "PG-100", "9", ""],  # the same pair again
            ]
        ),
        kind="SPECIAL_PRICES",
    )
    assert job["counts"]["new"] == 2 and job["counts"]["error"] == 3, job["errors"]
    assert messages(job) == [
        "Row 4, column “Shop”: No shop has the mobile number or code 9999999999.",
        "Row 5, column “Product code”: No product has the code NOPE.",
        "Row 6, column “Product code”: This shop and product are also in row 2. List each pair "
        "once.",
    ]
    free = [w for c in job["changes"] for w in c["warnings"]]
    assert free == [
        "A price of 0 makes this product free for this shop. Free-goods schemes are "
        "not supported yet."
    ]
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        prices = {r.product.code: r.price for r in RetailerPrice.objects.filter(retailer=shop)}
    assert prices == {"PG-100": D("8.75"), "MG-70": D("0.00")}
    assert AuditLog.objects.filter(action="pricing.retailer_price_changed").count() == 2

    again = upload(
        owner, run, _file([header, [shop.code, "PG-100", "8.50", ""]]), kind="SPECIAL_PRICES"
    )
    assert "choose “Add new and update existing”" in messages(again)[0]
    update = upload(
        owner,
        run,
        _file([header, [shop.code, "PG-100", "8.50", ""]]),
        mode="ADD_OR_UPDATE",
        kind="SPECIAL_PRICES",
    )
    [change] = update["changes"]
    assert (change["changes"], change["highlight"]) == (
        {"Special price": ["8.75", "8.50"]},
        ["Special price"],
    )

    exported = _export(owner, "retailer-prices/export/")
    assert exported[0][:4] == header
    assert sorted(r[:3] for r in exported[1:]) == [
        [shop.code, "MG-70", "0.00"],
        [shop.code, "PG-100", "8.75"],
    ]
    round_trip = upload(owner, run, _file(exported), mode="ADD_OR_UPDATE", kind="SPECIAL_PRICES")
    assert round_trip["counts"]["unchanged"] == 2 and round_trip["counts"]["changes"] == 0

    other = _client(tenant_b)
    assert _export(other, "retailer-prices/export/") == [exported[0]]
    assert _client(tenant_a, "WAREHOUSE").get(f"{API}/retailer-prices/export/").status_code == 403


# --- Price-list prices ---------------------------------------------------------------------------


@covers("price-list-items-export")
def test_price_list_prices_import_and_export(tenant_a, tenant_b, world, owner, run):
    header = ["Price list", "Product code", "Price"]
    job = upload(
        owner,
        run,
        _file(
            [
                header,
                ["gold", "PG-100", "9.00"],
                ["Silver", "PG-100", "9.00"],
                ["Gold", "MG-70", "0"],
            ]
        ),
        kind="PRICE_LIST_ITEMS",
    )
    assert messages(job) == [
        "Row 3, column “Price list”: No price list is called Silver. Your price lists: Gold."
    ]
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        items = dict(PriceListItem.objects.values_list("product__code", "price"))
    assert items == {"PG-100": D("9.00"), "MG-70": D("0.00")}

    update = upload(
        owner,
        run,
        _file([header, ["Gold", "PG-100", "9.40"], ["Gold", "MG-70", "0"]]),
        mode="ADD_OR_UPDATE",
        kind="PRICE_LIST_ITEMS",
    )
    assert update["counts"] == {
        "total": 2,
        "new": 0,
        "update": 1,
        "unchanged": 1,
        "error": 0,
        "changes": 1,
    }
    assert update["changes"][0]["highlight"] == ["Price"]
    exported = _export(owner, "price-lists/items/export/")
    assert exported[0][:3] == header and len(exported) == 3
    assert _export(_client(tenant_b), "price-lists/items/export/") == [exported[0]]


# --- Discount rules ------------------------------------------------------------------------------

RULE_HEADER = [
    "Rule name",
    "Product code",
    "Brand",
    "Category",
    "Price list",
    "Shop",
    "Type",
    "Discount",
    "From quantity",
    "Valid from",
    "Valid to",
    "Active",
]


@covers("discount-rules-export")
def test_discount_rules_import(tenant_a, tenant_b, world, owner, run):
    shop = world["shop"]
    job = upload(
        owner,
        run,
        _file(
            [
                RULE_HEADER,
                ["Biscuits bulk", "", "", "Food > Biscuits", "", "", "%", "2", "12", "", "", ""],
                ["Biscuits bulk", "", "", "Food > Biscuits", "", "", "%", "5", "24", "", "", ""],
                [
                    "Parle week",
                    "",
                    "Parle",
                    "",
                    "Gold",
                    "",
                    "%",
                    "3",
                    "",
                    "01-10-2026",
                    "15-10-2026",
                    "",
                ],
                [
                    "Ganesh deal",
                    "PG-100",
                    "",
                    "",
                    "",
                    "98765 00001",
                    "₹",
                    "1.50",
                    "",
                    "",
                    "",
                    "Yes",
                ],
                ["Bad mix", "", "", "", "", "", "%", "2", "12", "", "", ""],
                ["Bad mix", "", "", "", "", "", "%", "3", "", "", "", ""],
                ["Two targets", "PG-100", "Parle", "", "", "", "%", "2", "", "", "", ""],
                ["Too much", "", "", "", "", "", "%", "120", "", "", "", ""],
                ["Mismatch", "", "", "", "", "", "%", "2", "10", "", "", ""],
                ["Mismatch", "", "", "", "Gold", "", "%", "4", "20", "", "", ""],
            ]
        ),
        kind="DISCOUNT_RULES",
    )
    text = messages(job)
    assert "Row 7, column “From quantity”: Give every row of a rule with slabs a quantity." in text
    assert "Row 8, column “Brand”: Fill only one of product code, brand or category." in text
    assert "Row 9, column “Discount”: Enter a percentage above 0 and up to 100." in text
    assert (
        "Row 11, column “Price list”: Rows of the rule “Mismatch” must match row 10 here." in text
    )
    assert any(
        m.startswith("Row 10, column “Rule name”: Row 11 of this rule has a problem") for m in text
    )
    assert job["counts"]["new"] == 4, job["errors"]  # 2 slab rows + 2 single-row rules
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        bulk = DiscountRule.objects.get(name="Biscuits bulk")
        assert (bulk.scope_type, bulk.category_id) == ("CATEGORY", world["biscuits"].pk)
        assert [(s.min_qty, s.value) for s in bulk.slabs.order_by("min_qty")] == [
            (D("12.000"), D("2.00")),
            (D("24.000"), D("5.00")),
        ]
        week = DiscountRule.objects.get(name="Parle week")
        assert (week.audience_type, week.price_list_id, week.valid_from, week.valid_to) == (
            "PRICE_LIST",
            world["gold"].pk,
            date(2026, 10, 1),
            date(2026, 10, 15),
        )
        deal = DiscountRule.objects.get(name="Ganesh deal")
        assert (deal.discount_type, deal.value, deal.retailer_id, deal.product_id) == (
            "FLAT_PER_UNIT",
            D("1.50"),
            shop.pk,
            world["products"]["PG-100"].pk,
        )
    assert AuditLog.objects.filter(action="pricing.discount_rule_created").count() == 3

    # Update: the name identifies the rule; slabs are replaced; changes are previewed.
    again = upload(
        owner,
        run,
        _file(
            [
                RULE_HEADER,
                ["Biscuits bulk", "", "", "Food > Biscuits", "", "", "%", "3", "12", "", "", ""],
                ["Biscuits bulk", "", "", "Food > Biscuits", "", "", "%", "6", "36", "", "", ""],
            ]
        ),
        mode="ADD_OR_UPDATE",
        kind="DISCOUNT_RULES",
    )
    first = again["changes"][0]
    assert first["changes"] == {"Discount": ["12+: 2%; 24+: 5%", "12+: 3%; 36+: 6%"]}
    assert first["highlight"] == ["Discount"]
    commit(owner, run, again)
    with tenant_context(tenant_a.pk):
        bulk = DiscountRule.objects.get(name="Biscuits bulk")
        assert [(s.min_qty, s.value) for s in bulk.slabs.order_by("min_qty")] == [
            (D("12.000"), D("3.00")),
            (D("36.000"), D("6.00")),
        ]
        DiscountRule.objects.create(
            name="Parle week",
            discount_type="PERCENT",
            value=D("1"),
            scope_type="ALL",
            audience_type="ALL",
        )
    twice = upload(
        owner,
        run,
        _file([RULE_HEADER, ["Parle week", "", "Parle", "", "Gold", "", "%", "4", "", "", "", ""]]),
        mode="ADD_OR_UPDATE",
        kind="DISCOUNT_RULES",
    )
    assert messages(twice)[0].startswith("Row 2, column “Rule name”: 2 rules are called")
    exists = upload(
        owner,
        run,
        _file([RULE_HEADER, ["Ganesh deal", "", "", "", "", "", "%", "4", "", "", "", ""]]),
        kind="DISCOUNT_RULES",
    )
    assert "already exists" in messages(exists)[0]

    exported = _export(owner, "discount-rules/export/")
    assert exported[0] == RULE_HEADER
    bulk_rows = [r for r in exported if r[0] == "Biscuits bulk"]
    assert [(r[3], r[6], r[7], r[8]) for r in bulk_rows] == [
        ("Food > Biscuits", "%", "3", "12"),
        ("Food > Biscuits", "%", "6", "36"),
    ]
    assert _export(_client(tenant_b), "discount-rules/export/") == [RULE_HEADER]


def test_pricing_imports_need_pricing_manage(tenant_a, world, run):
    sales = _client(tenant_a, "SALES")  # pricing.view only
    for kind in ("SPECIAL_PRICES", "PRICE_LIST_ITEMS", "DISCOUNT_RULES"):
        response = run(
            sales.post,
            f"{API}/imports/",
            {"kind": kind, "mode": "ADD_ONLY", "file": _file([["x"]])},
            format="multipart",
        )
        assert response.status_code == 403, kind
    assert sales.get(f"{API}/discount-rules/export/").status_code == 200


def test_rule_round_trip_is_unchanged(tenant_a, world, owner, run):
    upload_job = upload(
        owner,
        run,
        _file(
            [
                RULE_HEADER,
                ["Bulk", "", "", "Food > Biscuits", "", "", "%", "2", "12", "", "", ""],
                ["Bulk", "", "", "Food > Biscuits", "", "", "%", "5", "24", "", "", ""],
                [
                    "Deal",
                    "PG-100",
                    "",
                    "",
                    "",
                    "98765 00001",
                    "₹",
                    "1.5",
                    "",
                    "01-10-2026",
                    "",
                    "No",
                ],
            ]
        ),
        kind="DISCOUNT_RULES",
    )
    commit(owner, run, upload_job)
    exported = _export(owner, "discount-rules/export/")
    again = upload(owner, run, _file(exported), mode="ADD_OR_UPDATE", kind="DISCOUNT_RULES")
    assert again["counts"]["changes"] == 0 and again["counts"]["error"] == 0, again["errors"]
