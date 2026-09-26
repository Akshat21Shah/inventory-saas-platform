"""Retailer import (ADR-035): matched by mobile; the number and sign-in never change on update;
new shops get their welcome message; credit columns need the credit permission."""

import io
from decimal import Decimal

import pytest
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounts.adapters.sms import MockSmsSender
from apps.accounts.models import User
from apps.accounts.tests.factories import make_staff_in
from apps.accounts.tokens import issue_tokens
from apps.retailers.models import Retailer
from common.storage import InMemoryStorage
from common.tenancy import tenant_context
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"


@pytest.fixture(autouse=True)
def _clean():
    cache.clear()
    InMemoryStorage.objects.clear()
    MockSmsSender.outbox.clear()
    yield
    cache.clear()
    InMemoryStorage.objects.clear()
    MockSmsSender.outbox.clear()


@pytest.fixture
def run(django_capture_on_commit_callbacks):
    def _run(fn, *args, **kwargs):
        with django_capture_on_commit_callbacks(execute=True):
            return fn(*args, **kwargs)

    return _run


def _client(tenant, role="OWNER", user=None):
    user = user or make_staff_in(tenant, role)
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {issue_tokens(user, tenant.pk).access}")
    client.defaults["HTTP_X_FORWARDED_HOST"] = f"{tenant.slug}.localhost"
    return client


def xlsx(rows):
    book = Workbook()
    sheet = book.worksheets[0]
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("retailers.xlsx", buffer.getvalue())


def upload(client, run, file, mode="ADD_ONLY"):
    response = run(
        client.post,
        f"{API}/imports/",
        {"kind": "RETAILERS", "mode": mode, "file": file},
        format="multipart",
    )
    assert response.status_code == 201, response.json()
    return client.get(f"{API}/imports/{response.json()['id']}/").json()


def commit(client, run, job):
    assert run(client.post, f"{API}/imports/{job['id']}/commit/").status_code == 200
    return client.get(f"{API}/imports/{job['id']}/").json()


def messages(job):
    return [m for row in job["errors"] for m in row["messages"]]


HEADER = [
    "Mobile",
    "Shop name",
    "Owner name",
    "GSTIN",
    "State",
    "Address",
    "City",
    "PIN code",
    "Salesperson email",
    "Credit limit",
    "Payment days",
    "Language",
]


def test_import_new_shops_with_messy_values(tenant_a, run):
    owner = _client(tenant_a)
    sales = make_staff_in(tenant_a, "SALES", email="ravi@alpha.example.com")
    job = upload(
        owner,
        run,
        xlsx(
            [
                HEADER,
                [
                    "+91 98765-00001 ",
                    "Ganesh Kirana",
                    "Ganesh",
                    "",
                    "Maharashtra",
                    "12 Station Rd",
                    "Pune",
                    411001,
                    "RAVI@alpha.example.com",
                    "₹50,000",
                    "15",
                    "Marathi",
                ],
                [
                    9876500002,
                    "Om Stores",
                    "",
                    "27AAGFK7315R1ZP",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "hindi",
                ],
                [None] * 12,
                ["09876500003", "Laxmi General", "", "", "27", "", "", "", "", "0", "", ""],
            ]
        ),
    )
    assert job["counts"]["new"] == 3 and job["counts"]["error"] == 0, job["errors"]
    done = commit(owner, run, job)
    assert done["counts"]["applied"] == 3
    with tenant_context(tenant_a.pk):
        ganesh = Retailer.objects.get(mobile="+919876500001")
        assert (
            ganesh.credit_limit,
            ganesh.payment_terms_days,
            ganesh.salesperson_id,
            ganesh.preferred_language,
        ) == (Decimal("50000.00"), 15, sales.pk, "mr")
        assert ganesh.addresses.get().pincode == "411001"
        om = Retailer.objects.get(mobile="+919876500002")
        assert (om.state_id, om.pan, om.preferred_language) == ("27", "AAGFK7315R", "hi")
        assert Retailer.objects.get(mobile="+919876500003").credit_limit == Decimal("0.00")
    assert User.objects.filter(tenant=tenant_a, user_type="RETAILER").count() == 3
    assert [m.template for m in MockSmsSender.outbox] == ["retailer_welcome"] * 3


def test_errors_name_the_row_and_column(tenant_a, run):
    owner = _client(tenant_a)
    job = upload(
        owner,
        run,
        xlsx(
            [
                HEADER,
                ["12345", "Short number", "", "", "27", "", "", "", "", "", "", ""],
                ["9876500001", "", "", "", "27", "", "", "", "", "", "", ""],
                ["9876500002", "No state", "", "", "", "", "", "", "", "", "", ""],
                ["9876500003", "Bad GSTIN", "", "27AAGFK7315R1ZA", "", "", "", "", "", "", "", ""],
                [
                    "9876500004",
                    "Wrong state",
                    "",
                    "27AAGFK7315R1ZP",
                    "Karnataka",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ],
                ["98765 00004", "Twice", "", "", "27", "", "", "", "", "", "", ""],
                ["9876500005", "Half address", "", "", "27", "Road", "", "", "", "", "", ""],
                [
                    "9876500006",
                    "Unknown staff",
                    "",
                    "",
                    "27",
                    "",
                    "",
                    "",
                    "nobody@x.com",
                    "",
                    "",
                    "",
                ],
            ]
        ),
    )
    text = messages(job)
    assert "Row 2, column “Mobile”: Enter a 10-digit Indian mobile number." in text
    assert "Row 3, column “Shop name”: Needed for a new shop." in text
    assert "Row 4, column “State”: Needed when the shop has no GSTIN." in text
    assert any(m.startswith("Row 5, column “GSTIN”:") and "O/0" in m for m in text)
    assert any(m.startswith("Row 6, column “State”: The state must match") for m in text)
    assert any(m.startswith("Row 7, column “Mobile”: 98765 00004 is also in row 6") for m in text)
    assert any(m.startswith("Row 8, column “Address”: For an address") for m in text)
    assert any(m.startswith("Row 9, column “Salesperson email”:") for m in text)


def test_update_never_changes_the_mobile_or_login_and_credit_needs_permission(tenant_a, run):
    owner = _client(tenant_a)
    commit(
        owner,
        run,
        upload(
            owner,
            run,
            xlsx(
                [
                    HEADER,
                    ["9876500001", "Ganesh Kirana", "Ganesh", "", "27", "", "", "", "", "", "", ""],
                ]
            ),
        ),
    )
    MockSmsSender.outbox.clear()
    job = upload(
        owner,
        run,
        xlsx(
            [
                ["Mobile", "Shop name", "Owner name", "Credit limit"],
                ["9876500001", "Ganesh Super Mart", "", "10000"],
            ]
        ),
        mode="ADD_OR_UPDATE",
    )
    [change] = job["changes"]
    assert change["changes"] == {
        "Shop name": ["Ganesh Kirana", "Ganesh Super Mart"],
        "Credit limit": ["", "10000.00"],
    }
    assert change["highlight"] == ["Credit limit"]
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        shop = Retailer.objects.get()
        assert (shop.shop_name, shop.owner_name, shop.mobile) == (
            "Ganesh Super Mart",
            "Ganesh",
            "+919876500001",
        )
    assert MockSmsSender.outbox == []  # no welcome for an existing shop

    sales = _client(tenant_a, "SALES")  # retailers.manage without credit.manage
    denied = upload(
        sales, run, xlsx([["Mobile", "Credit limit"], ["9876500001", "1"]]), mode="ADD_OR_UPDATE"
    )
    assert any("You can't set credit limits" in m for m in messages(denied))
    add_only = upload(
        owner, run, xlsx([HEADER, ["9876500001", "X", "", "", "27", "", "", "", "", "", "", ""]])
    )
    assert any("A shop with mobile 9876500001 already exists" in m for m in messages(add_only))


@covers("retailers-export")
def test_export_round_trip_and_isolation(tenant_a, tenant_b, run):
    owner = _client(tenant_a)
    commit(
        owner,
        run,
        upload(
            owner,
            run,
            xlsx(
                [
                    HEADER,
                    [
                        "9876500001",
                        "Ganesh Kirana",
                        "Ganesh",
                        "",
                        "27",
                        "12 Road",
                        "Pune",
                        "411001",
                        "",
                        "",
                        "",
                        "Marathi",
                    ],
                ]
            ),
        ),
    )
    exported = owner.get(f"{API}/retailers/export/", {"file_type": "xlsx"})
    sheet = load_workbook(io.BytesIO(exported.content)).worksheets[0]
    assert [c.value for c in sheet[2]][:2] == ["9876500001", "Ganesh Kirana"]
    again = upload(
        owner, run, SimpleUploadedFile("export.xlsx", exported.content), mode="ADD_OR_UPDATE"
    )
    assert again["counts"]["unchanged"] == 1 and again["counts"]["error"] == 0, again["errors"]
    other = _client(tenant_b).get(f"{API}/retailers/export/", {"file_type": "csv"})
    assert b"9876500001" not in other.content
    warehouse = _client(tenant_a, "WAREHOUSE")  # no retailers.view
    assert warehouse.get(f"{API}/retailers/export/").status_code == 403
