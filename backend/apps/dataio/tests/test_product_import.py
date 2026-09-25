"""Product import (ADR-035) with the files people really send: title rows, merged and padded
headers, blank rows and columns, trailing spaces, numbers stored as text, ₹ and commas, duplicate
codes, missing columns, wrong GST rates, bad HSN codes, and CSVs saved from Excel in different
encodings. Every error names the row (as numbered in the file) and the column, in plain words."""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.catalog import selectors
from apps.catalog.models import Brand, Category, Product
from apps.dataio.models import ImportJob
from apps.dataio.tests.helpers import _client, commit, csv_file, messages, upload, xlsx
from common.tenancy import tenant_context
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"
HEADER = [
    "Product code",
    "Product name",
    "Unit",
    "HSN code",
    "GST rate",
    "Price",
    "MRP",
    "Category",
    "Brand",
]


def _rate(product_id):
    row = selectors.tax_rate_on(product_id)
    assert row is not None
    return row.gst_rate


# --- The happy path ----------------------------------------------------------------------------


@covers("imports", "import-detail", "import-commit")
def test_clean_file_validates_then_imports_with_new_brands_and_categories(owner, tenant_a, run):
    job = upload(
        owner,
        run,
        xlsx(
            [
                HEADER,
                [
                    "PG-100",
                    "Parle-G 100g",
                    "PCS",
                    "1905",
                    "5",
                    "9.00",
                    "10.00",
                    "Food > Biscuits",
                    "Parle",
                ],
                [
                    "MG-70",
                    "Maggi 70g",
                    "PCS",
                    "1902",
                    "5",
                    "12.50",
                    "14.00",
                    "Food > Noodles",
                    "Nestle",
                ],
                ["TS-1", "Tata Salt 1kg", "PKT", "2501", "0", "25", "28", "Food", "Tata"],
            ]
        ),
    )
    assert job["status"] == "VALIDATED"
    assert job["counts"] == {
        "total": 3,
        "new": 3,
        "update": 0,
        "unchanged": 0,
        "error": 0,
        "changes": 3,
    }
    with tenant_context(tenant_a.pk):
        assert Product.objects.count() == 0  # nothing is saved before the distributor confirms

    done = commit(owner, run, job)
    assert done["status"] == "COMMITTED" and done["counts"]["applied"] == 3
    with tenant_context(tenant_a.pk):
        product = Product.objects.get(code="PG-100")
        assert product.base_price == Decimal("9.00") and product.brand.name == "Parle"
        assert product.category.name == "Biscuits" and product.category.parent.name == "Food"
        assert _rate(product.pk) == Decimal("5")
        assert Brand.objects.count() == 3 and Category.objects.count() == 3
    assert AuditLog.objects.filter(action="catalog.product_created").count() == 3
    assert AuditLog.objects.filter(action="import.committed").exists()


# --- Messy files -------------------------------------------------------------------------------


def test_messy_excel_file_is_understood(owner, tenant_a, run):
    rows = [
        ["Sharma Distributors — price list September", None, None, None, None, None, None],
        [None] * 7,
        [
            "  product CODE * ",
            "Product Name*",
            None,
            "UOM",
            "HSN/SAC",
            "GST %",
            "Selling Price (Rs)",
            "M.R.P.",
        ],
        ["PG-100 ", "  Parle-G   100g ", None, " pcs", "1905", "5%", "₹1,234.50", "Rs. 1,299/-"],
        [None] * 8,
        ["'77", "Milk powder", None, "PCS", 402, "5", "1,00,000", None],
        [None, None, None, None, None, None, None, None],
        ["MG-70", "Maggi", None, "PCS", "19023010", " 18 ", "12.5", "14"],
    ]
    file = xlsx(rows, merge=["A1:G1"])
    job = upload(owner, run, file)
    assert job["status"] == "VALIDATED", job
    assert job["counts"]["new"] == 3 and job["counts"]["error"] == 0, job["errors"]
    by_key = {row["key"]: row for row in job["changes"]}
    assert set(by_key) == {"PG-100", "'77", "MG-70"} or set(by_key) == {"PG-100", "77", "MG-70"}
    notes = [w for row in job["changes"] for w in row["warnings"]]
    assert any("read as 0402" in w for w in notes)  # Excel dropped the leading zero
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        pg = Product.objects.get(code="PG-100")
        assert (pg.name, pg.base_price, pg.mrp) == (
            "Parle-G 100g",
            Decimal("1234.50"),
            Decimal("1299.00"),
        )
        milk = Product.objects.exclude(code__in=["PG-100", "MG-70"]).get()
        assert (milk.hsn_code, milk.base_price) == ("0402", Decimal("100000.00"))
        assert _rate(Product.objects.get(code="MG-70").pk) == Decimal("18")


def test_merged_header_cells_and_blank_columns(owner, tenant_a, run):
    """A header cell merged over two columns ("Product name" over B:C) and a blank column."""
    rows = [
        ["Product code", "Product name", None, None, "Unit", "HSN code", "GST rate", "Price"],
        ["SP-1", "Surf Excel 1kg", None, None, "PCS", "3402", "18", "220"],
        ["SP-2", "Surf Excel 500g", "stray note", None, "PCS", "3402", "18", "115"],
    ]
    job = upload(owner, run, xlsx(rows, merge=["B1:C1"]))
    assert job["counts"]["new"] == 2 and job["counts"]["error"] == 0, job["errors"]
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        assert Product.objects.get(code="SP-2").name == "Surf Excel 500g"


def test_row_numbers_and_columns_in_messages_match_the_file(owner, run):
    job = upload(
        owner,
        run,
        xlsx(
            [
                ["My products"],
                HEADER,
                ["A-1", "Good row", "PCS", "1905", "5", "10", "12", "", ""],
                [
                    "A-2",
                    "Bad GST",
                    "PCS",
                    "1905",
                    "12",
                    "10",
                    "12",
                    "",
                    "",
                ],  # 12% is no longer in use
                ["A-3", "Bad HSN", "PCS", "19A5", "5", "10", "12", "", ""],
                ["A-1", "Duplicate", "PCS", "1905", "5", "10", "12", "", ""],
                ["A-5", "Bad price", "PCS", "1905", "5", "ten", "12", "", ""],
                ["A-6", "Bad unit", "TRAY", "1905", "5", "10", "", "", ""],
                ["A-7", "", "PCS", "1905", "5", "10", "", "", ""],
                ["A-8", "Negative", "PCS", "1905", "5", "-3", "", "", ""],
                ["A-9", "Too deep", "PCS", "1905", "7", "10", "", "A > B > C > D", ""],
            ]
        ),
    )
    assert job["counts"]["new"] == 1 and job["counts"]["error"] == 8
    text = messages(job)
    assert any(
        m.startswith("Row 4, column “GST rate”: 12% isn't a GST rate in use. Use one of")
        for m in text
    ), text
    assert "Row 5, column “HSN code”: 19A5 isn't a valid HSN code. Use 4 to 8 digits." in text
    duplicate = "Row 6, column “Product code”: A-1 is also in row 3. List each product only once."
    assert duplicate in text
    assert any(m.startswith("Row 7, column “Price”: ten isn't a number") for m in text)
    assert any(m.startswith("Row 8, column “Unit”: TRAY isn't one of your units") for m in text)
    assert "Row 9, column “Product name”: Needed for a new product." in text
    assert "Row 10, column “Price”: Can't be negative." in text
    assert any("Row 11, column “GST rate”: 7% isn't a GST rate in use" in m for m in text)
    assert "Row 11, column “Category”: Categories can be at most 3 levels deep." in text

    report = load_workbook(io.BytesIO(owner.get(f"{API}/imports/{job['id']}/report/").content))
    statuses = [row[2] for row in report.worksheets[0].iter_rows(min_row=2, values_only=True)]
    assert statuses.count("Error") == 8 and statuses.count("New") == 1


def test_missing_required_columns_fail_the_file_in_plain_words(owner, run):
    job = upload(owner, run, xlsx([["Product code", "Product name", "Price"], ["A", "B", "1"]]))
    assert job["status"] == "FAILED"
    assert job["problem"] == (
        "The file is missing these columns: Unit, HSN code, GST rate. Download the template to "
        "see the expected columns."
    )


@pytest.mark.parametrize(
    ("encoding", "bom", "delimiter"),
    [
        ("utf-8", b"", ","),
        ("utf-8", b"\xef\xbb\xbf", ","),
        ("utf-16", b"", "\t"),
        ("cp1252", b"", ";"),
    ],
    ids=["utf8", "utf8-bom", "utf16-excel-unicode-text", "windows-1252-semicolon"],
)
def test_csv_from_excel_in_any_common_encoding(owner, tenant_a, run, encoding, bom, delimiter):
    rows = [HEADER, ["CF-1", "Café Coffee 50g", "PCS", "0901", "5", "Rs. 1,200.00", "", "", ""]]
    job = upload(owner, run, csv_file(rows, encoding=encoding, delimiter=delimiter, bom=bom))
    assert job["status"] == "VALIDATED", job
    assert job["counts"]["new"] == 1, job["errors"]
    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        product = Product.objects.get(code="CF-1")
        assert (product.name, product.base_price) == ("Café Coffee 50g", Decimal("1200.00"))


@pytest.mark.parametrize(
    ("file", "problem"),
    [
        (
            SimpleUploadedFile("old.xls", b"\xd0\xcf\x11\xe0 old binary"),
            "Old .xls files aren't supported",
        ),
        (SimpleUploadedFile("notes.pdf", b"%PDF-1.4"), "Upload an Excel (.xlsx) or CSV file."),
        (
            SimpleUploadedFile("broken.xlsx", b"PK\x03\x04broken"),
            "This isn't a readable Excel file",
        ),
        (
            SimpleUploadedFile("no-header.csv", b"a,b,c\n1,2,3\n"),
            "We couldn't find the column names",
        ),
    ],
)
def test_unreadable_files_get_one_clear_message(owner, run, file, problem):
    job = upload(owner, run, file)
    assert job["status"] == "FAILED" and problem in job["problem"]


def test_empty_file_is_refused_on_upload(owner):
    response = owner.post(
        f"{API}/imports/",
        {"kind": "PRODUCTS", "mode": "ADD_ONLY", "file": SimpleUploadedFile("empty.csv", b"")},
        format="multipart",
    )
    assert response.json()["error"]["details"]["fields"]["file"] == ["The file is empty."]


def test_mode_must_be_chosen(owner, run):
    response = owner.post(
        f"{API}/imports/", {"kind": "PRODUCTS", "file": xlsx([HEADER])}, format="multipart"
    )
    assert "Add new only" in response.json()["error"]["details"]["fields"]["mode"][0]


# --- Updating existing products ----------------------------------------------------------------


def _seed(owner, run):
    upload_job = upload(
        owner,
        run,
        xlsx(
            [
                HEADER,
                ["PG-100", "Parle-G 100g", "PCS", "1905", "5", "9.00", "10.00", "Food", "Parle"],
                ["MG-70", "Maggi 70g", "PCS", "1902", "5", "12.50", "14.00", "Food", "Nestle"],
            ]
        ),
    )
    commit(owner, run, upload_job)


def test_add_only_reports_existing_codes(owner, run):
    _seed(owner, run)
    job = upload(owner, run, xlsx([HEADER, ["pg-100", "X", "PCS", "1905", "5", "9", "", "", ""]]))
    assert messages(job) == [
        "Row 2, column “Product code”: A product with code pg-100 already exists. To change "
        "it, choose “Add new and update existing”."
    ]


def test_update_import_previews_old_and_new_values_and_only_changes_given_cells(
    owner, tenant_a, run
):
    _seed(owner, run)
    job = upload(
        owner,
        run,
        xlsx(
            [
                ["Product code", "Price", "MRP", "Product name", "Brand"],
                ["PG-100", "9.50", "", "Parle-G Gold 100g", ""],  # blank MRP and brand: no change
                ["MG-70", "12.50", "14", "", ""],  # same values: unchanged
                ["NEW-1", "5", "", "Needs more columns", ""],  # a new code in an update file
            ]
        ),
        mode="ADD_OR_UPDATE",
    )
    assert job["counts"] == {
        "total": 3,
        "new": 0,
        "update": 1,
        "unchanged": 1,
        "error": 1,
        "changes": 1,
    }
    [change] = job["changes"]
    assert change["key"] == "PG-100" and change["action"] == "UPDATE"
    assert change["changes"] == {
        "Price": ["9.00", "9.50"],
        "Product name": ["Parle-G 100g", "Parle-G Gold 100g"],
    }
    assert change["highlight"] == ["Price"]
    assert any("Needed for a new product" in m for m in messages(job))

    commit(owner, run, job)
    with tenant_context(tenant_a.pk):
        product = Product.objects.get(code="PG-100")
        assert (product.base_price, product.mrp, product.brand.name) == (
            Decimal("9.50"),
            Decimal("10.00"),
            "Parle",
        )
    price_change = AuditLog.objects.get(action="catalog.product_price_changed")
    assert price_change.changes == {"base_price": ["9.00", "9.50"]}  # audited like a manual edit


def test_update_preview_highlights_price_and_mrp_changes_only(owner, run):
    _seed(owner, run)
    job = upload(
        owner,
        run,
        xlsx(
            [
                ["Product code", "Price", "MRP", "Product name"],
                ["PG-100", "", "11.00", ""],  # MRP only
                ["MG-70", "13.00", "15.00", "Maggi 70g pack"],  # price, MRP and name
            ]
        ),
        mode="ADD_OR_UPDATE",
    )
    by_key = {row["key"]: row for row in job["changes"]}
    assert by_key["PG-100"]["changes"] == {"MRP": ["10.00", "11.00"]}
    assert by_key["PG-100"]["highlight"] == ["MRP"]
    assert by_key["MG-70"]["changes"]["Product name"] == ["Maggi 70g", "Maggi 70g pack"]
    assert sorted(by_key["MG-70"]["highlight"]) == ["MRP", "Price"]  # the name isn't highlighted
    assert job["counts"]["changes"] == 2


def test_update_import_never_changes_gst_silently(owner, run):
    _seed(owner, run)
    job = upload(
        owner, run, xlsx([["Product code", "GST rate"], ["PG-100", "18"]]), mode="ADD_OR_UPDATE"
    )
    assert any(
        "GST changes need a start date" in m and "current rate is 5%" in m for m in messages(job)
    )


def test_commit_rechecks_the_file(owner, tenant_a, run):
    row = ["LATE-1", "Late", "PCS", "1905", "5", "9", "", "", ""]
    job = upload(owner, run, xlsx([HEADER, row]))
    owner.post(
        f"{API}/products/",
        {
            "code": "LATE-1",
            "name": "Made by hand",
            "unit": str(_pcs(tenant_a)),
            "hsn_code": "1905",
            "base_price": "1",
            "gst_rate": "5",
        },
        format="json",
    )
    done = commit(owner, run, job)
    assert done["counts"]["applied"] == 0
    with tenant_context(tenant_a.pk):
        assert Product.objects.get(code="LATE-1").name == "Made by hand"


def _pcs(tenant):
    from apps.catalog.models import Unit

    with tenant_context(tenant.pk):
        return Unit.objects.get(code="PCS").pk


# --- Templates, export round trip, permissions, isolation --------------------------------------


@covers("import-template", "products-export")
def test_template_and_export_round_trip(owner, run):
    template = load_workbook(io.BytesIO(owner.get(f"{API}/imports/templates/products/").content))
    header = [c.value for c in template.worksheets[0][1]]
    assert header[:3] == ["Product code *", "Product name *", "Unit *"]
    assert template.worksheets[1].title == "How to fill"

    _seed(owner, run)
    exported = owner.get(f"{API}/products/export/", {"file_type": "xlsx"})
    assert exported.status_code == 200, exported.content[:300]
    assert exported["Content-Disposition"].endswith('products.xlsx"')
    again = upload(
        owner, run, SimpleUploadedFile("export.xlsx", exported.content), mode="ADD_OR_UPDATE"
    )
    assert again["counts"]["unchanged"] == 2 and again["counts"]["error"] == 0, again["errors"]
    csv_export = owner.get(f"{API}/products/export/", {"file_type": "csv"})
    assert csv_export.content.startswith(b"\xef\xbb\xbf") and b"PG-100" in csv_export.content


@covers("import-report")
def test_imports_are_isolated_and_need_the_right_permission(owner, tenant_a, tenant_b, run):
    job = upload(owner, run, xlsx([HEADER, ["A-1", "X", "PCS", "1905", "5", "9", "", "", ""]]))
    other = _client(tenant_b)
    assert other.get(f"{API}/imports/").json()["results"] == []
    assert other.get(f"{API}/imports/{job['id']}/").status_code == 404
    assert other.get(f"{API}/imports/{job['id']}/report/").status_code == 404
    assert other.post(f"{API}/imports/{job['id']}/commit/").status_code == 404

    sales = _client(tenant_a, "SALES")  # products.view only
    denied = sales.post(
        f"{API}/imports/",
        {"kind": "PRODUCTS", "mode": "ADD_ONLY", "file": xlsx([HEADER])},
        format="multipart",
    )
    assert denied.status_code == 403
    assert sales.get(f"{API}/imports/templates/products/").status_code == 403
    assert sales.get(f"{API}/imports/{job['id']}/").status_code == 404
    assert sales.get(f"{API}/products/export/").status_code == 200  # viewing products is enough
    warehouse = _client(tenant_a, "WAREHOUSE")
    assert warehouse.get(f"{API}/imports/").json()["results"] == []
    assert ImportJob.objects.unscoped().count() == 1
