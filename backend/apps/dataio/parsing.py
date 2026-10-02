"""Reading spreadsheets people actually send (ADR-035): xlsx or csv, any common encoding, a title
row or blank rows above the header, merged or padded header cells, blank rows and columns, trailing
spaces, numbers stored as text, prices with commas and ₹. Nothing here knows about products or
retailers: it returns the header and the non-blank rows, each with its row number in the file.
"""

import codecs
import csv
import io
import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from django.utils.translation import gettext as _

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
HEADER_SEARCH_ROWS = 15


class FileProblem(Exception):
    """The file as a whole can't be read (wrong type, empty, no header...). Plain message."""


@dataclass(frozen=True)
class Row:
    number: int  # as the user sees it: the sheet row or the CSV line
    values: dict[str, str]  # canonical column → cleaned text ("" = blank)


@dataclass(frozen=True)
class Sheet:
    columns: list[str]  # canonical names found, in file order
    ignored: list[str]  # header cells that match no known column
    rows: list[Row]


def normalize_header(text: str) -> str:
    """ "GST Rate (%) *" → "gst rate"; "M.R.P." → "mrp"; "HSN/SAC" → "hsn sac"."""
    text = text.lower().replace("*", " ").replace("%", " ")
    text = re.sub(r"\(.*?\)", " ", text)
    text = text.replace(".", "")
    text = re.sub(r"[^0-9a-z]+", " ", text)
    return " ".join(text.split())


def clean_text(value: Any) -> str:
    """Cell → text: trims and collapses spaces; whole-number floats lose ``.0`` (codes typed as
    numbers); dates become ISO dates."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(str(value)), "f")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    # str.split() also splits on non-breaking spaces (common in text copied from the web).
    return " ".join(str(value).split())


_MONEY_NOISE = re.compile(r"(?i)^(?:rs\.?|inr|₹)\s*|\s*(?:rs\.?|inr|/-)$|₹|,")


def parse_decimal(text: str, *, places: int) -> Decimal:
    """ "₹1,23,456.50" / "Rs. 99/-" / " 18% " / "1,000" → Decimal. Raises ``ValueError``."""
    cleaned = _MONEY_NOISE.sub("", text.strip()).replace(" ", "").rstrip("%")
    if not cleaned:
        raise ValueError("empty")
    try:
        number = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(text) from exc
    if not number.is_finite():
        raise ValueError(text)
    quantum = Decimal(1).scaleb(-places)
    if number != number.quantize(quantum):
        raise ValueError(f"more than {places} decimal places")
    return number.quantize(quantum)


def parse_bool(text: str) -> bool:
    lowered = text.strip().lower()
    if lowered in {"yes", "y", "true", "1", "on", "active", "show"}:
        return True
    if lowered in {"no", "n", "false", "0", "off", "inactive", "hide"}:
        return False
    raise ValueError(text)


_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y")


def parse_date(text: str) -> date:
    """ "2026-10-01" (Excel date cells) / "01-10-2026" / "1/10/2026" → date (day first, as in
    India). Raises ``ValueError``."""
    cleaned = text.strip().split(" ")[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def split_list(text: str) -> list[str]:
    """ "a, b; c" → ["a", "b", "c"]."""
    return [part.strip() for part in re.split(r"[,;|]", text) if part.strip()]


def decode_csv(data: bytes) -> str:
    """UTF-8 (with or without BOM), UTF-16 (Excel "Unicode text") or Windows-1252 (Excel
    "CSV" on Windows)."""
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _csv_grid(data: bytes) -> list[list[Any]]:
    text = decode_csv(data)
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _xlsx_grid(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        # Not read-only: merged cells are unmerged (value in the first cell, None elsewhere).
        book = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise FileProblem(
            _("This isn't a readable Excel file. Save it as .xlsx and try again.")
        ) from exc
    sheet = book.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def read_grid(file_name: str, data: bytes) -> list[list[Any]]:
    if not data:
        raise FileProblem(_("The file is empty."))
    if len(data) > MAX_FILE_BYTES:
        raise FileProblem(_("The file is larger than 10 MB. Split it into smaller files."))
    name = file_name.lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        return _xlsx_grid(data)
    if name.endswith((".csv", ".txt")):
        return _csv_grid(data)
    if name.endswith(".xls"):
        raise FileProblem(_("Old .xls files aren't supported. Save the file as .xlsx or .csv."))
    raise FileProblem(_("Upload an Excel (.xlsx) or CSV file."))


def read_sheet(file_name: str, data: bytes, synonyms: dict[str, str]) -> Sheet:
    """Find the header (the first row, among the top rows, naming the most known columns) and
    return the rows below it. ``synonyms`` maps normalized header text → canonical column."""
    grid = read_grid(file_name, data)
    best_index, best_hits = -1, 0
    for index, row in enumerate(grid[:HEADER_SEARCH_ROWS]):
        hits = len({synonyms[h] for c in row if (h := normalize_header(clean_text(c))) in synonyms})
        if hits > best_hits:
            best_index, best_hits = index, hits
    if best_index < 0:
        raise FileProblem(
            _(
                "We couldn't find the column names. Use the template, or put the column names in "
                "the first row."
            )
        )
    header = grid[best_index]
    positions: dict[int, str] = {}
    ignored: list[str] = []
    for position, cell in enumerate(header):
        text = clean_text(cell)
        if not text:
            continue  # blank or the tail of a merged header cell
        canonical = synonyms.get(normalize_header(text))
        if canonical is None or canonical in positions.values():
            ignored.append(text)
            continue
        positions[position] = canonical
    rows: list[Row] = []
    for offset, raw in enumerate(grid[best_index + 1 :], start=best_index + 2):
        values = {
            name: clean_text(raw[pos]) if pos < len(raw) else "" for pos, name in positions.items()
        }
        if not any(values.values()):
            continue  # blank rows anywhere are skipped
        rows.append(Row(number=offset, values=values))
        if len(rows) > MAX_ROWS:
            raise FileProblem(
                _("The file has more than %(max_rows)s rows. Split it into smaller files.")
                % {"max_rows": format(MAX_ROWS, ",")}
            )
    if not rows:
        raise FileProblem(_("The file has column names but no rows to import."))
    return Sheet(columns=list(positions.values()), ignored=ignored, rows=rows)


def synonyms_for(columns: Iterable[tuple[str, Iterable[str]]]) -> dict[str, str]:
    """``[(canonical, [header spellings...])]`` → normalized spelling → canonical."""
    mapping: dict[str, str] = {}
    for canonical, spellings in columns:
        for spelling in (canonical, *spellings):
            mapping[normalize_header(spelling)] = canonical
    return mapping
