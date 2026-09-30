"""The GST summary for filing (ADR-050 item 9): each kind of sale lands in its GSTR-1 section, in
the official template's layout (sheet names, summary rows, header row, date and place-of-supply
formats, UQC codes), for a month or a quarter."""

import io
from datetime import date, timedelta
from decimal import Decimal as D
from typing import Any

import pytest
from openpyxl import load_workbook

from apps.accounts.tests.factories import make_staff_in
from apps.billing import credit_notes
from apps.billing.credit_notes import ReturnLine
from apps.billing.models import Invoice
from apps.billing.tests.helpers import ship_invoice
from apps.catalog.models import Product, ProductTaxRate, Unit
from apps.inventory import services as inventory
from apps.orders.tests.helpers import add_stock, client_for, make_shop
from apps.platform.tests.factories import make_gstin
from apps.reports.definitions.gst import check_period, gst_date
from apps.reports.models import ReportRun
from apps.retailers.services import AddressInput, create_retailer
from common.dates import today_ist
from common.storage import get_storage
from common.tenancy import tenant_context

pytestmark = pytest.mark.django_db
API = "/api/v1"
HEADERS = {
    "b2b,sez,de": [
        "GSTIN/UIN of Recipient",
        "Receiver Name",
        "Invoice Number",
        "Invoice date",
        "Invoice Value",
        "Place Of Supply",
        "Reverse Charge",
        "Applicable % of Tax Rate",
        "Invoice Type",
        "E-Commerce GSTIN",
        "Rate",
        "Taxable Value",
        "Cess Amount",
    ],
    "hsn(b2c)": [
        "HSN",
        "Description",
        "UQC",
        "Total Quantity",
        "Total Value",
        "Rate",
        "Taxable Value",
        "Integrated Tax Amount",
        "Central Tax Amount",
        "State/UT Tax Amount",
        "Cess Amount",
    ],
}


def product(tenant: Any, code: str, rate: str, price: str, hsn: str) -> Product:
    with tenant_context(tenant.pk):
        made: Product = Product.objects.create(
            code=code,
            name=f"Product {code}",
            unit=Unit.objects.get(code="PCS"),
            hsn_code=hsn,
            base_price=D(price),
        )
        ProductTaxRate.objects.create(product=made, gst_rate=D(rate), effective_from=today_ist())
        inventory.ensure_levels_exist([made.pk], inventory.default_warehouse())
    add_stock(tenant, made, "100")
    return made


def shop(tenant: Any, phone: str, name: str, state: str, gstin: str = "") -> Any:
    with tenant_context(tenant.pk):
        return create_retailer(
            shop_name=name,
            phone=phone,
            gstin=gstin,
            billing=AddressInput("1 Road", "City", "560001" if state == "29" else "411001", state),
            send_welcome=False,
        )


@pytest.fixture
def world(tenant_a, tenant_b):
    owner = make_staff_in(tenant_a, "OWNER")
    biscuit = product(tenant_a, "BIS", "18", "100", "1905")
    laptop = product(tenant_a, "LAP", "18", "30000", "8471")
    rice = product(tenant_a, "RICE", "0", "50", "1006")  # nil-rated
    local_b2b = shop(tenant_a, "9876500091", "Local Wholesale", "27", make_gstin(9101, "27"))
    far_b2b = shop(tenant_a, "9876500092", "Kaveri Traders", "29", make_gstin(9102, "29"))
    far_b2c = shop(tenant_a, "9876500093", "Bengaluru Retail", "29")
    local_b2c = make_shop(tenant_a, "9876500094", shop_name="Ganesh Kirana")
    b2b = ship_invoice(tenant_a, local_b2b, owner, (biscuit, "10"), (rice, "2"))
    b2b_far = ship_invoice(tenant_a, far_b2b, owner, (biscuit, "1"))
    b2cl = ship_invoice(tenant_a, far_b2c, owner, (laptop, "4"))  # ₹1,41,600 > ₹1 lakh
    b2cs = ship_invoice(tenant_a, local_b2c, owner, (biscuit, "3"))
    cancelled = ship_invoice(tenant_a, local_b2c, owner, (biscuit, "1"))
    with tenant_context(tenant_a.pk):
        for invoice, qty in ((b2b, "2"), (b2cl, "1"), (b2cs, "1")):
            line = invoice.lines.get(product__code__in=("BIS", "LAP"))
            credit_notes.issue_return(
                invoice.pk, [ReturnLine(line.pk, D(qty))], reason="DAMAGED", by=owner
            )
        Invoice.objects.filter(pk=cancelled.pk).update(status="CANCELLED")  # its IRN cancelled
    first = today_ist().replace(day=1)
    last = (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    period = f"date_from={first.isoformat()}&date_to={last.isoformat()}"
    return {
        "t": tenant_a,
        "owner": client_for(tenant_a, owner),
        "sales": client_for(tenant_a, make_staff_in(tenant_a, "SALES")),
        "invoices": {"b2b": b2b, "far": b2b_far, "b2cl": b2cl, "b2cs": b2cs, "gone": cancelled},
        "period": period,
    }


def test_only_a_month_or_a_quarter():
    assert check_period({"date_from": date(2026, 9, 1), "date_to": date(2026, 9, 30)}) == {}
    assert check_period({"date_from": date(2026, 7, 1), "date_to": date(2026, 9, 30)}) == {}
    assert check_period({"date_from": date(2027, 1, 1), "date_to": date(2027, 3, 31)}) == {}
    assert check_period({"date_from": date(2026, 8, 1), "date_to": date(2026, 10, 31)}) == {
        "date_to": ["Choose a whole month or a whole quarter."]
    }
    assert "date_from" in check_period(
        {"date_from": date(2026, 9, 2), "date_to": date(2026, 9, 30)}
    )
    assert gst_date(date(2026, 9, 1)) == "01-Sep-2026"


def test_the_summary_on_screen(world):
    response = world["owner"].get(f"{API}/reports/gst_summary/?{world['period']}")
    assert response.status_code == 200, response.json()
    body = response.json()
    rows = {r["section"]: r for r in body["rows"]}
    assert rows["B2B invoices (4A)"]["documents"] == 2
    assert (
        rows["B2C large invoices (5)"]["documents"],
        rows["B2C large invoices (5)"]["taxable"],
    ) == (
        1,
        "120000.00",
    )
    assert rows["B2C others (7)"]["taxable"] == "200.00"  # 300 - 100 returned, netted
    assert rows["Credit notes, registered (9B)"]["documents"] == 1
    assert rows["Credit notes, unregistered (9B)"]["documents"] == 1
    assert body["notes"][1] == (
        "Nil-rated (0%) supplies, not in these sections (Table 8): "
        "Intra-State supplies to registered persons ₹100.00."
    )
    bad = world["owner"].get(f"{API}/reports/gst_summary/?date_from=2026-09-01&date_to=2026-09-15")
    assert bad.status_code == 400
    assert world["sales"].get(f"{API}/reports/gst_summary/?{world['period']}").status_code == 403


def test_the_workbook_follows_the_gstr1_template(world, django_capture_on_commit_callbacks):
    with django_capture_on_commit_callbacks(execute=True):
        queued = world["owner"].post(
            f"{API}/reports/gst_summary/export/",
            {"filters": dict(p.split("=") for p in world["period"].split("&"))},
            format="json",
        )
    assert queued.status_code == 202  # always in the background
    with tenant_context(world["t"].pk):
        run = ReportRun.objects.get(pk=queued.json()["id"])
    assert run.status == "READY"
    book = load_workbook(io.BytesIO(get_storage().get(run.file_key)))
    assert book.sheetnames == [
        "b2b,sez,de",
        "b2cl",
        "b2cs",
        "cdnr",
        "cdnur",
        "hsn(b2b)",
        "hsn(b2c)",
        "docs",
        "About",
    ]
    b2b = [list(r) for r in book["b2b,sez,de"].values]
    assert b2b[0][0] == "Summary For B2B, SEZ, DE (4A, 4B, 6B, 6C)" and b2b[0][12] == "HELP"
    assert b2b[3] == HEADERS["b2b,sez,de"]
    invoices = world["invoices"]
    local = next(r for r in b2b[4:] if r[2] == invoices["b2b"].number)
    assert local[3] == gst_date(invoices["b2b"].invoice_date)
    assert (local[5], local[6], local[8], local[10], local[11]) == (
        "27-Maharashtra",
        "N",
        "Regular B2B",
        18,
        1000,
    )
    far = next(r for r in b2b[4:] if r[2] == invoices["far"].number)
    assert far[5] == "29-Karnataka"
    assert invoices["gone"].number not in [r[2] for r in b2b[4:]]
    assert len(b2b) == 4 + 2  # the 0% line is not a B2B row
    b2cl = [list(r) for r in book["b2cl"].values]
    assert b2cl[1][0] == "No. of Invoices" and not any(b2cl[1][1:])  # as the template
    assert (b2cl[4][0], b2cl[4][3], b2cl[4][5], b2cl[4][6]) == (
        invoices["b2cl"].number,
        "29-Karnataka",
        18,
        120000,
    )
    b2cs = [list(r) for r in book["b2cs"].values]
    assert b2cs[4][:6] == ["OE", "27-Maharashtra", None, 18, 200, 0]
    cdnr = [list(r) for r in book["cdnr"].values]
    assert (cdnr[4][4], cdnr[4][7], cdnr[4][11]) == ("C", "Regular B2B", 200)
    cdnur = [list(r) for r in book["cdnur"].values]
    assert (cdnur[4][0], cdnur[4][3], cdnur[4][8]) == ("B2CL", "C", 30000)
    hsn = [list(r) for r in book["hsn(b2c)"].values]
    assert hsn[3] == HEADERS["hsn(b2c)"]
    laptop = next(r for r in hsn[4:] if r[0] == "8471")
    assert (laptop[1], laptop[2], laptop[3], laptop[5], laptop[6]) == (
        None,
        "PCS-PIECES",
        3,
        18,
        90000,
    )
    rice = next(r for r in [list(x) for x in book["hsn(b2b)"].values][4:] if r[0] == "1006")
    assert (rice[5], rice[6]) == (0, 100)  # nil-rated lines are in the HSN summary
    docs = [list(r) for r in book["docs"].values]
    invoice_row = next(r for r in docs[4:] if r[0] == "Invoices for outward supply")
    assert (invoice_row[3], invoice_row[4]) == (5, 1)  # the cancelled one counted as cancelled
    credit_row = next(r for r in docs[4:] if r[0] == "Credit Note")
    assert credit_row[3] == 3
