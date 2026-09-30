"""The report framework (ADR-050), with a small report over shops registered for the test: the
permission, own shops only where the sales-visibility rule applies, cost columns left out
without costs.view (rows, totals and files), filters, paging and totals, exports at once or in
the background with "Report ready", the requester's own exports only, expiry, and isolation."""

import io
from datetime import timedelta
from decimal import Decimal
from typing import Any

import pytest
from django.db.models import F, Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.tests.factories import make_staff_in
from apps.notifications.models import Notification
from apps.orders.tests.helpers import client_for, make_shop, settings
from apps.reports import services
from apps.reports.models import ReportRun
from apps.reports.registry import (
    PERIOD,
    REGISTRY,
    Column,
    Context,
    Filter,
    FilterKind,
    Group,
    Kind,
    Report,
    register,
)
from apps.retailers.models import Retailer
from common.dates import today_ist
from common.permissions import AnyOf
from common.storage import get_storage
from common.tenancy import tenant_context
from common.testing.isolation import covers

pytestmark = pytest.mark.django_db
API = "/api/v1"
CODE = "test_shops"


def _rows(ctx: Context) -> Any:
    rows = ctx.shops(Retailer.objects.all(), path="")
    if ctx.params.get("status"):
        rows = rows.filter(status=ctx.params["status"])
    return rows.values("shop_name", "credit_limit", retailer_id=F("id")).order_by("shop_name")


def _totals(ctx: Context) -> dict[str, Any]:
    found: dict[str, Any] = _rows(ctx).aggregate(credit_limit=Sum("credit_limit"))
    return found


@pytest.fixture
def report():
    r = register(
        Report(
            code=CODE,
            title="Shops (test)",
            group=Group.SALES,
            permission=AnyOf(("reports.sales", "reports.sales_own")),
            full="reports.sales",
            columns=(
                Column("shop_name", "Shop", width=30),
                Column("credit_limit", "Credit limit", Kind.MONEY, cost=True, total=True),
            ),
            filters=(*PERIOD, Filter("status", "Status", FilterKind.CHOICE, choices=("ACTIVE",))),
            rows=_rows,
            totals=_totals,
            pdf=True,
            max_days=31,
        )
    )
    yield r
    REGISTRY.pop(CODE, None)


@pytest.fixture
def world(tenant_a, tenant_b, django_capture_on_commit_callbacks):
    owner = make_staff_in(tenant_a, "OWNER")
    sales = make_staff_in(tenant_a, "SALES")
    warehouse = make_staff_in(tenant_a, "WAREHOUSE")
    shops = [make_shop(tenant_a, f"98765001{n:02d}", shop_name=f"Shop {n}") for n in range(1, 4)]
    with tenant_context(tenant_a.pk):
        for n, shop in enumerate(shops, 1):
            Retailer.objects.filter(pk=shop.pk).update(credit_limit=Decimal(n * 1000))
        Retailer.objects.filter(pk=shops[0].pk).update(salesperson=sales)
    return {
        "t": tenant_a,
        "tb": tenant_b,
        "owner": client_for(tenant_a, owner),
        "owner_user": owner,
        "sales": client_for(tenant_a, sales),
        "sales_user": sales,
        "warehouse": client_for(tenant_a, warehouse),
        "other": client_for(tenant_b, make_staff_in(tenant_b, "OWNER")),
        "shops": shops,
        "run": lambda: django_capture_on_commit_callbacks(execute=True),
    }


@covers("reports")
def test_the_catalogue_offers_what_the_user_may_open(world, report):
    mine = {r["code"]: r for r in world["owner"].get(f"{API}/reports/").json()}
    assert [c["key"] for c in mine[CODE]["columns"]] == ["shop_name", "credit_limit"]
    assert mine[CODE]["filters"][0]["default"] == today_ist().replace(day=1).isoformat()
    sales = {r["code"]: r for r in world["sales"].get(f"{API}/reports/").json()}
    assert [c["key"] for c in sales[CODE]["columns"]] == ["shop_name"]  # no costs.view
    assert CODE not in {r["code"] for r in world["warehouse"].get(f"{API}/reports/").json()}


@covers("report")
def test_a_page_with_totals_filters_and_checks(world, report):
    body = world["owner"].get(f"{API}/reports/{CODE}/?page_size=2").json()
    assert body["count"] == 3 and len(body["rows"]) == 2
    assert body["rows"][0] == {
        "shop_name": "Shop 1",
        "credit_limit": "1000.00",
        "retailer_id": str(world["shops"][0].pk),
    }
    assert body["totals"] == {"credit_limit": "6000.00"}  # the whole report, not the page
    second = world["owner"].get(f"{API}/reports/{CODE}/?page_size=2&page=2").json()
    assert [r["shop_name"] for r in second["rows"]] == ["Shop 3"]
    errors = (
        world["owner"]
        .get(f"{API}/reports/{CODE}/?date_from=2026-01-01&date_to=2026-03-01&status=NOPE")
        .json()["error"]["details"]["fields"]
    )
    assert errors == {
        "status": ["Choose one of: ACTIVE."],
    }
    too_long = world["owner"].get(f"{API}/reports/{CODE}/?date_from=2026-01-01&date_to=2026-03-01")
    assert too_long.json()["error"]["details"]["fields"] == {"date_to": ["Choose at most 31 days."]}
    backwards = world["owner"].get(f"{API}/reports/{CODE}/?date_from=2026-03-02&date_to=2026-03-01")
    assert backwards.status_code == 400
    assert world["warehouse"].get(f"{API}/reports/{CODE}/").status_code == 403
    assert world["owner"].get(f"{API}/reports/nothing/").status_code == 404


def test_sales_staff_see_their_own_shops_when_the_distributor_says_so(world, report):
    everyone = world["sales"].get(f"{API}/reports/{CODE}/").json()
    assert (everyone["count"], everyone["own_shops"]) == (3, False)
    settings(world["t"], orders__sales_visibility="ASSIGNED_RETAILERS")
    own = world["sales"].get(f"{API}/reports/{CODE}/").json()
    assert (own["count"], own["own_shops"], own["totals"]) == (1, True, {})  # no cost total
    assert own["rows"] == [{"shop_name": "Shop 1", "retailer_id": str(world["shops"][0].pk)}]
    assert world["owner"].get(f"{API}/reports/{CODE}/").json()["count"] == 3


@covers("report-export")
def test_a_small_export_comes_at_once_without_what_the_user_may_not_see(world, report):
    settings(world["t"], orders__sales_visibility="ASSIGNED_RETAILERS")
    got = world["sales"].post(f"{API}/reports/{CODE}/export/", {"format": "XLSX"}, format="json")
    assert got.status_code == 200
    assert got["Content-Disposition"].startswith('attachment; filename="test-shops-')
    book = load_workbook(io.BytesIO(got.content))
    rows = list(book["Shops (test)"].values)
    assert rows == [("Shop",), ("Shop 1",)]  # own shops; no cost column, no totals row
    about: dict[Any, Any] = dict(book["About"].values)  # type: ignore[arg-type]
    assert about["Shops"] == "Only the shops assigned to you"
    assert about["Costs"].startswith("Cost columns are left out")
    owner = world["owner"].post(f"{API}/reports/{CODE}/export/", {"format": "XLSX"}, format="json")
    full = list(load_workbook(io.BytesIO(owner.content))["Shops (test)"].values)
    assert full[-1] == ("Total", Decimal("6000.00"))
    pdf = world["owner"].post(f"{API}/reports/{CODE}/export/", {"format": "PDF"}, format="json")
    assert (pdf.status_code, pdf["Content-Type"]) == (200, "application/pdf")
    bad = world["owner"].post(f"{API}/reports/{CODE}/export/", {"format": "CSV"}, format="json")
    assert bad.status_code == 400


@covers("report-runs", "report-run")
def test_a_background_export_brings_report_ready(world, report):
    register(Report(**{**report.__dict__, "code": "test_shops_big", "background_only": True}))
    try:
        with world["run"]():
            queued = world["owner"].post(
                f"{API}/reports/test_shops_big/export/", {"format": "XLSX"}, format="json"
            )
        assert queued.status_code == 202
        run_id = queued.json()["id"]
        shown = world["owner"].get(f"{API}/report-runs/{run_id}/").json()
        assert (shown["status"], shown["row_count"]) == ("READY", 3)
        assert shown["download_url"]
        with tenant_context(world["t"].pk):
            run = ReportRun.objects.get(pk=run_id)
            ready = Notification.objects.get(event_code="report.ready")
        assert get_storage().get(run.file_key)[:2] == b"PK"  # an xlsx (zip) file
        assert (ready.recipient_id, ready.channel, ready.title) == (
            world["owner_user"].pk,
            "IN_APP",
            "Report ready: Shops (test)",
        )
        assert [r["id"] for r in world["owner"].get(f"{API}/report-runs/").json()["results"]] == [
            run_id
        ]
        # Only the requester sees it; nobody in another business.
        assert world["sales"].get(f"{API}/report-runs/{run_id}/").status_code == 404
        assert world["sales"].get(f"{API}/report-runs/").json()["results"] == []
        assert world["other"].get(f"{API}/report-runs/{run_id}/").status_code == 404
        assert world["other"].get(f"{API}/report-runs/").json()["results"] == []
        with world["run"]():  # their own shops, their own export
            theirs = (
                world["other"]
                .post(f"{API}/reports/test_shops_big/export/", {}, format="json")
                .json()
            )
        listed = world["other"].get(f"{API}/report-runs/").json()["results"]
        assert ([r["id"] for r in listed], theirs["row_count"]) == ([theirs["id"]], None)
    finally:
        REGISTRY.pop("test_shops_big", None)


def test_more_rows_than_the_limit_go_to_the_background(world, report, monkeypatch):
    monkeypatch.setattr(services, "get_platform_setting", lambda key: 2)  # 3 shops > 2
    with world["run"]():
        queued = world["owner"].post(f"{API}/reports/{CODE}/export/", {}, format="json")
    assert (queued.status_code, queued.json()["status"]) == (202, "QUEUED")  # as it was sent


def test_the_file_is_deleted_when_its_link_runs_out(world, report):
    big = register(Report(**{**report.__dict__, "code": "test_shops_old", "background_only": True}))
    try:
        with world["run"]():
            run_id = (
                world["owner"]
                .post(f"{API}/reports/test_shops_old/export/", {}, format="json")
                .json()["id"]
            )
        with tenant_context(world["t"].pk):
            run = ReportRun.objects.get(pk=run_id)
            key = run.file_key
            ReportRun.objects.filter(pk=run_id).update(expires_at=timezone.now() - timedelta(1))
        assert services.expire_due() == 1
        shown = world["owner"].get(f"{API}/report-runs/{run_id}/").json()
        assert (shown["status"], shown["download_url"]) == ("EXPIRED", None)
        with pytest.raises(Exception):  # noqa: B017 - the storage says "missing" its own way
            get_storage().get(key)
    finally:
        REGISTRY.pop(big.code, None)


def test_a_person_no_longer_allowed_gets_no_file(world, report):
    big = register(
        Report(**{**report.__dict__, "code": "test_shops_gone", "background_only": True})
    )
    try:
        with world["run"]():
            run_id = (
                world["sales"]
                .post(f"{API}/reports/test_shops_gone/export/", {}, format="json")
                .json()["id"]
            )
            # (made at once by the captured callbacks: take the permission away and make again)
        with tenant_context(world["t"].pk):
            ReportRun.objects.filter(pk=run_id).update(status="QUEUED", file_key="")
            world["sales_user"].memberships.update(is_active=False)
            assert services.make_run(run_id) == "refused"
            run = ReportRun.objects.get(pk=run_id)
            failed = Notification.objects.filter(event_code="report.failed").count()
        assert (run.status, run.error) == ("FAILED", "You can no longer open this report.")
        assert failed == 0  # an inactive person gets nothing
    finally:
        REGISTRY.pop(big.code, None)


def test_report_messages_are_not_configurable(world, report):
    matrix = world["owner"].get(f"{API}/notification-rules/").json()
    assert "report.ready" not in {e["code"] for e in matrix["events"]}
    assert "REQUESTER" not in matrix["recipients"]
    changed = world["owner"].put(
        f"{API}/notification-rules/report.ready/",
        {"rules": [{"recipient": "OWNERS", "channels": ["IN_APP"]}]},
        format="json",
    )
    assert changed.status_code == 404
    requester = world["owner"].put(
        f"{API}/notification-rules/order.placed/",
        {"rules": [{"recipient": "REQUESTER", "channels": ["IN_APP"]}]},
        format="json",
    )
    assert requester.status_code == 400
