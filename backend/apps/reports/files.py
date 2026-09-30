"""Report files (ADR-050): Excel for every report, PDF where offered. Excel is written in
openpyxl's streaming mode so large exports use little memory; an "About" sheet says what the
file contains (filters, who asked and when, own shops only, notes)."""

from __future__ import annotations

import io
import re
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.billing.pdf import get_renderer
from apps.billing.templatetags.documents import qty, rupees
from apps.reports.engine import columns, row_out
from apps.reports.registry import Column, Context, Kind, Report, Sheet
from common.dates import to_ist

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
BOLD = Font(bold=True)
FORMATS = {Kind.MONEY: "#,##0.00", Kind.QTY: "#,##0.###", Kind.DATE: "DD-MM-YYYY"}


def file_name(report: Report, ctx: Context, extension: str) -> str:
    start, end = ctx.params.get("date_from"), ctx.params.get("date_to")
    period = f"-{start:%Y%m%d}-{end:%Y%m%d}" if isinstance(start, date) and end else ""
    return f"{report.code.replace('_', '-')}{period}.{extension}"


def _sheet_title(title: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "-", title)[:31] or "Report"


def _excel_value(column: Column, value: Any) -> Any:
    if value is None:
        return None
    if column.kind in (Kind.MONEY, Kind.QTY, Kind.PERCENT):
        return Decimal(value)
    if column.kind == Kind.INT:
        return int(value)
    if column.kind == Kind.DATE and isinstance(value, datetime):
        return to_ist(value).date()
    return value if isinstance(value, date | int | float | Decimal) else str(value)


def _cells(
    sheet: Any, cols: tuple[Column, ...], row: dict[str, Any], bold: bool = False
) -> list[Any]:
    out = []
    for c in cols:
        cell = WriteOnlyCell(sheet, value=_excel_value(c, row.get(c.key)))
        if c.kind in FORMATS:
            cell.number_format = FORMATS[c.kind]
        if bold:
            cell.font = BOLD
        out.append(cell)
    return out


def _write(book: Workbook, sheet: Sheet, totals: dict[str, Any] | None = None) -> int:
    ws = book.create_sheet(title=_sheet_title(sheet.title))
    for index, c in enumerate(sheet.columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = c.width
    header = []
    for c in sheet.columns:
        cell = WriteOnlyCell(ws, value=c.label)
        cell.font = BOLD
        header.append(cell)
    ws.append(header)
    count = 0
    for row in sheet.rows:
        ws.append(_cells(ws, sheet.columns, row))
        count += 1
    if totals:
        first = sheet.columns[0].key
        ws.append(_cells(ws, sheet.columns, {**totals, first: "Total"}, bold=True))
    return count


def about(
    report: Report, ctx: Context, described: list[tuple[str, str]], by: str
) -> list[tuple[str, str]]:
    lines: list[tuple[str, str]] = [
        ("Report", report.title),
        *described,
        ("Made", f"{to_ist(timezone.now()):%d-%m-%Y %H:%M} (IST) for {by}"),
    ]
    if ctx.scope.own_shops:
        lines.append(("Shops", "Only the shops assigned to you"))
    if not ctx.scope.costs and any(c.cost for c in report.columns):
        lines.append(("Costs", "Cost columns are left out (they need permission to see costs)"))
    lines += [("Note", note) for note in (report.notes(ctx) if report.notes else [])]
    return lines


def excel(report: Report, ctx: Context, about_lines: list[tuple[str, str]]) -> tuple[bytes, int]:
    """The workbook and how many data rows it holds."""
    book = Workbook(write_only=True)
    cols = columns(report, ctx.scope)
    if report.sheets is not None:
        rows = 0
        for sheet in report.sheets(ctx):
            visible = tuple(c for c in sheet.columns if ctx.scope.costs or not c.cost)
            rows += _write(book, Sheet(sheet.title, visible, sheet.rows))
    else:
        rows = _write(
            book, Sheet(report.title, cols, _iterate(report.rows(ctx))), _totals(report, ctx, cols)
        )
    info = book.create_sheet(title="About")
    for label, value in about_lines:
        info.append([label, value])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), rows


def _totals(report: Report, ctx: Context, cols: tuple[Column, ...]) -> dict[str, Any] | None:
    """The totals row, when a column the user sees has a total."""
    if report.totals is None or not any(c.total for c in cols):
        return None
    return report.totals(ctx)


def _iterate(rows: Any) -> Iterable[dict[str, Any]]:
    """Large querysets are read in chunks, not all at once."""
    found: Iterable[dict[str, Any]] = (
        rows.iterator(chunk_size=2000) if hasattr(rows, "iterator") else rows
    )
    return found


def _pdf_value(column: Column, value: Any) -> str:
    if value is None or value == "":
        return ""
    if column.kind == Kind.MONEY:
        return rupees(Decimal(value))
    if column.kind == Kind.QTY:
        return qty(Decimal(value))
    if column.kind == Kind.PERCENT:
        return f"{Decimal(value):.1f}%"
    if column.kind == Kind.DATE:
        day = value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
        return f"{day:%d-%m-%Y}"
    return str(value)


def pdf(report: Report, ctx: Context, about_lines: list[tuple[str, str]]) -> tuple[bytes, int]:
    cols = columns(report, ctx.scope)
    rows = [row_out(r, cols) for r in _iterate(report.rows(ctx))]
    totals = _totals(report, ctx, cols)
    numeric = [c.kind in (Kind.MONEY, Kind.QTY, Kind.INT, Kind.PERCENT) for c in cols]
    total_row = None
    if totals:
        total_row = [
            ("Total" if i == 0 else _pdf_value(c, totals.get(c.key)) if c.total else "", n)
            for i, (c, n) in enumerate(zip(cols, numeric, strict=True))
        ]
    html = render_to_string(
        "reports/report.html",
        {
            "title": report.title,
            "about": about_lines,
            "headers": list(zip([c.label for c in cols], numeric, strict=True)),
            "rows": [
                list(zip([_pdf_value(c, r.get(c.key)) for c in cols], numeric, strict=True))
                for r in rows
            ],
            "totals": total_row,
        },
    )
    return get_renderer().render(html), len(rows)
