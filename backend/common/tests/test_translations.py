"""Server messages in the person's language (ADR-060, 11a.4): the catalogs are complete and up to
date with the code; field errors, the framework's own messages, import results and account
emails come in Hindi or Marathi; English is unchanged."""

import ast
import io

import pytest
from django.core import mail
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import translation
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.accounts.tests.factories import make_staff_in
from apps.orders.tests.helpers import client_for
from common import languages, translations
from common.errors import DomainError, InvalidFields
from common.storage import InMemoryStorage

pytestmark = pytest.mark.django_db
API = "/api/v1"


@pytest.fixture(autouse=True)
def _clean():
    cache.clear()
    InMemoryStorage.objects.clear()
    yield
    cache.clear()
    InMemoryStorage.objects.clear()


def test_every_language_has_every_message_with_the_same_placeholders():
    found = [p for lang in translations.translated_languages() for p in translations.problems(lang)]
    assert found == []  # run `make messages` (and translate what it lists)
    assert len(translations.extract()) > 650


def test_messages_are_written_so_they_can_be_translated():
    """gettext must get the English as it is: an f-string or a variable can't be looked up."""
    wrong = []
    for path in translations._source_files():
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if not isinstance(node, ast.Call) or not node.args:
                continue
            func = node.func
            name = func.attr if isinstance(func, ast.Attribute) else getattr(func, "id", "")
            if name in translations.KEYWORDS:
                message = node.args[1 if name.startswith("pgettext") else 0]
                if not (isinstance(message, ast.Constant) and isinstance(message.value, str)):
                    wrong.append(f"{path.relative_to(translations.BACKEND)}:{node.lineno}")
    assert wrong == []


def test_placeholders_are_compared_both_styles():
    assert translations.placeholders("Enter %(low)s to %(high)s days, 5%%.") == ["", "high", "low"]
    assert translations.placeholders("At most {max_length} characters.") == ["max_length"]
    assert translations.placeholders("Only {{ variable }} placeholders.") == []


@pytest.mark.usefixtures("every_language")
def test_field_errors_come_in_the_language_the_web_asks_for(tenant_a):
    owner = client_for(tenant_a, make_staff_in(tenant_a, "OWNER"))
    body = {"shop_name": "Shree Kirana", "mobile": "12", "state_code": "27"}
    english = owner.post(f"{API}/retailers/", body, format="json").json()
    hindi = owner.post(f"{API}/retailers/", body, format="json", HTTP_ACCEPT_LANGUAGE="hi")
    marathi = owner.post(f"{API}/retailers/", body, format="json", HTTP_ACCEPT_LANGUAGE="mr")
    fields = english["error"]["details"]["fields"]
    assert fields["mobile"] == ["Enter a 10-digit Indian mobile number."]
    assert hindi.json()["error"]["details"]["fields"]["mobile"] == [
        "10 अंकों का भारतीय मोबाइल नंबर डालें।"
    ]
    assert marathi.json()["error"]["details"]["fields"]["mobile"] == [
        "10 अंकी भारतीय मोबाईल नंबर भरा."
    ]
    # The framework's own messages too (DRF has no Hindi or Marathi of its own).
    missing = owner.post(f"{API}/retailers/", {}, format="json", HTTP_ACCEPT_LANGUAGE="mr")
    assert missing.json()["error"]["details"]["fields"]["shop_name"] == ["हा रकाना आवश्यक आहे."]


def test_an_error_keeps_the_language_it_was_raised_in():
    """Stored errors (an import's rows, a failed job) don't change language when read later."""
    with translation.override("hi"):
        error = InvalidFields({"name": [translation.gettext("Enter a name.")]})
        default = DomainError()
    assert error.details == {"fields": {"name": ["नाम डालें।"]}}
    assert default.message == "अनुरोध पूरा नहीं हो सका।"
    assert isinstance(default.message, str)


def _sheet(rows: list[list[object]]) -> SimpleUploadedFile:
    book = Workbook()
    for row in rows:
        book.worksheets[0].append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("shops.xlsx", buffer.getvalue())


@pytest.mark.usefixtures("every_language")
def test_an_import_is_checked_in_the_language_of_whoever_uploaded_it(
    tenant_a, django_capture_on_commit_callbacks
):
    person = make_staff_in(tenant_a, "OWNER")
    User.objects.filter(pk=person.pk).update(preferred_language="mr")
    owner = client_for(tenant_a, person)
    with django_capture_on_commit_callbacks(execute=True):
        created = owner.post(
            f"{API}/imports/",
            {
                "kind": "RETAILERS",
                "mode": "ADD_ONLY",
                "file": _sheet([["Mobile", "Shop name", "State"], ["9876500191", "", "27"]]),
            },
            format="multipart",
        )
    job = owner.get(f"{API}/imports/{created.json()['id']}/").json()
    # Read back in English: the stored text stays in the uploader's language.
    assert job["errors"][0]["messages"] == ["ओळ 2, कॉलम “Shop name”: नवीन दुकानासाठी आवश्यक."]


@pytest.mark.usefixtures("every_language")
def test_account_emails_come_in_the_persons_language(tenant_a, django_capture_on_commit_callbacks):
    user = make_staff_in(tenant_a, "OWNER", email="asha@example.com")
    User.objects.filter(pk=user.pk).update(preferred_language="hi", full_name="Asha")
    with django_capture_on_commit_callbacks(execute=True):
        APIClient().post(
            f"{API}/auth/password/forgot/",
            {"email": "asha@example.com"},
            format="json",
            HTTP_X_FORWARDED_HOST="localhost",
        )
    email = mail.outbox[-1]
    assert email.subject == "अपना पासवर्ड बदलें"
    assert email.body.startswith("नमस्ते Asha,\n\nकिसी ने आपके खाते का पासवर्ड")
    assert "/reset-password/" in email.body


def test_without_a_review_hindi_isnt_used_for_a_persons_emails(tenant_a):
    """Until Hindi is on for this distributor, its people get English (ADR-060 item 12)."""
    user = make_staff_in(tenant_a, "OWNER")
    User.objects.filter(pk=user.pk).update(preferred_language="hi")
    user.refresh_from_db()
    with languages.speaking(user, tenant_a.pk):
        assert translation.get_language() == "en"
