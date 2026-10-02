import pyotp
import pytest
from django.core.management import CommandError, call_command

from apps.accounts import mfa
from apps.accounts.models import Membership, User
from apps.catalog.models import Category, Product, ProductImage
from apps.platform.models import Tenant
from apps.pricing.models import DiscountRule, PriceList, RetailerPrice
from apps.retailers.models import Retailer
from apps.shop.selectors import ShopFilters, priced, shop_products
from common import demo
from common.tenancy import tenant_context

pytestmark = pytest.mark.django_db


def test_seed_is_idempotent(settings):
    settings.DEBUG = True
    call_command("seed", "--no-photos")
    call_command("seed", "--no-photos")
    assert User.objects.filter(email="admin@platform.local", is_superuser=True).count() == 1
    assert set(Tenant.objects.values_list("slug", flat=True)) == {"sharma", "patel"}
    sharma = Tenant.objects.get(slug="sharma")
    with tenant_context(sharma.id):
        roles = set(Membership.objects.values_list("role__code", flat=True))
        assert roles == {"OWNER", "MANAGER", "SALES", "WAREHOUSE", "ACCOUNTS"}
        assert Retailer.objects.count() == 20
    owner = User.objects.get(email="owner@sharma.example.com")
    assert owner.check_password("staff-dev-password")
    # The shared shop phone has a login under both distributors (sign-in chooser).
    assert User.objects.filter(phone="+919876500000", user_type=User.UserType.RETAILER).count() == 2
    admin = User.objects.get(email="admin@platform.local")
    assert admin.totp_enabled
    assert mfa.matching_step(admin.totp_secret, pyotp.TOTP(admin.totp_secret).now()) is not None
    # Demo billing (Phase 5): run once, and the ledger reconciles.
    from apps.billing.models import CreditNote, Invoice
    from apps.ledger.tests.helpers import check_ledger
    from apps.payments.models import Payment, Refund

    with tenant_context(sharma.id):
        assert Invoice.objects.exists()
        assert Payment.objects.count() == 4
        assert Payment.objects.filter(handover_status="WITH_SALESMAN").count() == 1
        assert CreditNote.objects.count() == 1
        assert Refund.objects.count() == 1
    check_ledger(sharma)
    # Demo notifications (Phase 6): half the shops on WhatsApp, one announcement, once.
    from apps.notifications.models import Announcement

    with tenant_context(sharma.id):
        assert Retailer.objects.filter(whatsapp_opt_in=True).count() == 10
        assert Announcement.objects.count() == 1
    # Phase 7: e-invoices, e-way bills and online payments on for Sharma only, with the mocks.
    from apps.compliance.models import GstCredential
    from apps.payments.models import GatewayConfig
    from apps.platform.selectors import effective_features
    from apps.retailers.models import RetailerAddress

    patel = Tenant.objects.get(slug="patel")
    on = ("einvoice", "ewaybill", "payments")
    assert all(effective_features(sharma.id)[code] for code in on)
    assert not any(effective_features(patel.id)[code] for code in on)
    with tenant_context(sharma.id):
        assert GstCredential.objects.get().status == "VERIFIED"
        assert GatewayConfig.objects.get().provider == "MOCK"
        assert not RetailerAddress.objects.filter(distance_km__isnull=True).exists()
    with tenant_context(patel.id):
        assert not GstCredential.objects.exists() and not GatewayConfig.objects.exists()
    # Phase 9a: purchasing and stock planning for Sharma only, with orders in every state.
    from apps.planning.models import ProductStats, ReorderSuggestion
    from apps.purchasing.models import PurchaseOrder, Supplier, SupplierProduct

    nine = ("purchasing", "stock_planning")
    assert all(effective_features(sharma.id)[code] for code in nine)
    assert not any(effective_features(patel.id)[code] for code in nine)
    with tenant_context(sharma.id):
        assert Supplier.objects.count() == 3
        assert not SupplierProduct.objects.filter(is_preferred=True).count() < 1
        statuses = sorted(PurchaseOrder.objects.values_list("status", flat=True))
        assert statuses == ["DRAFT", "PARTLY_RECEIVED", "SENT", "SENT"]
        assert ProductStats.objects.exists() and ReorderSuggestion.objects.exists()
    with tenant_context(patel.id):
        assert not Supplier.objects.exists() and not ProductStats.objects.exists()
    # Phase 9b: shop activity for both; free-goods schemes for Sharma only.
    from apps.insights.models import Segment, ShopActivity, ShopContact
    from apps.pricing.models import FreeGoodsScheme

    assert effective_features(sharma.id)["free_goods"]
    assert not effective_features(patel.id)["free_goods"]
    with tenant_context(sharma.id):
        assert FreeGoodsScheme.objects.count() == 3
        assert ShopActivity.objects.count() == 20 and ShopContact.objects.count() == 1
        assert ShopActivity.objects.filter(segment=Segment.NEVER_ORDERED).exists()
    with tenant_context(patel.id):
        assert not FreeGoodsScheme.objects.exists() and ShopActivity.objects.exists()
    # Phase 9c: a return request waiting at Sharma.
    from apps.billing.models import ReturnRequest

    with tenant_context(sharma.id):
        assert list(ReturnRequest.objects.values_list("status", flat=True)) == ["REQUESTED"]


def test_seed_keeps_an_existing_admin_2fa_key(settings):
    settings.DEBUG = True
    call_command("seed", "--no-photos")
    admin = User.objects.get(email="admin@platform.local")
    admin.totp_secret = pyotp.random_base32()
    admin.save(update_fields=["totp_secret"])
    call_command("seed", "--no-photos")
    admin.refresh_from_db()
    assert admin.totp_secret != "DEVSEEDADMINTOTPKEYDEVSEEDADMIN2"


def test_seed_refuses_without_debug(settings):
    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command("seed", "--no-photos")


def test_seed_can_put_the_admin_back_on_the_dev_2fa_key(settings):
    settings.DEBUG = True
    call_command("seed", "--no-photos")
    admin = User.objects.get(email="admin@platform.local")
    admin.totp_secret = pyotp.random_base32()
    admin.save(update_fields=["totp_secret"])
    call_command("seed", "--no-photos", "--reset-admin-2fa")
    admin.refresh_from_db()
    assert admin.totp_secret == "DEVSEEDADMINTOTPKEYDEVSEEDADMIN2"


def test_seed_builds_a_catalog_each_shop_can_browse_at_its_own_prices(settings):
    settings.DEBUG = True
    call_command("seed", "--no-photos")
    sharma = Tenant.objects.get(slug="sharma")
    with tenant_context(sharma.id):
        assert Product.objects.count() == 200
        assert Product.objects.filter(brand__own_brand=True).count() == 20
        assert not Product.objects.filter(cost_price__isnull=True).exists()
        assert Category.objects.filter(level=3).exists()
        assert set(PriceList.objects.values_list("name", flat=True)) == {"Gold", "Wholesale"}
        assert DiscountRule.objects.count() == 5 and RetailerPrice.objects.count() == 5
        assert Retailer.objects.filter(status="BLOCKED").count() == 1
        gold_shop = Retailer.objects.filter(price_list__name="Gold").first()
        assert gold_shop is not None
        visible = list(shop_products(gold_shop, ShopFilters()))
        assert 150 < len(visible) < 200  # hidden and inactive items are left out
        sources = {r.price_source for _, r in priced(gold_shop, visible)}
        assert "PRICE_LIST" in sources and "BASE" in sources
    patel = Tenant.objects.get(slug="patel")
    with tenant_context(patel.id):
        assert not Product.objects.filter(code__startswith="SH-").exists()  # no mixing


def test_seed_photos_are_processed(settings, monkeypatch):
    settings.DEBUG = True
    monkeypatch.setattr(demo, "PRODUCTS_PER_TENANT", 3)
    call_command("seed")
    sharma = Tenant.objects.get(slug="sharma")
    with tenant_context(sharma.id):
        statuses = list(ProductImage.objects.values_list("status", flat=True))
    assert statuses == ["READY"] * 3


def test_e2e_workbook_is_an_importable_excel_file(settings, capsys):
    import base64
    import io

    from openpyxl import load_workbook

    settings.DEBUG = True
    call_command("e2e_workbook", "products", "--count", "3", "--prefix", "T1")
    book = load_workbook(io.BytesIO(base64.b64decode(capsys.readouterr().out)))
    rows = list(book.worksheets[0].iter_rows(values_only=True))
    assert rows[0][0] == "Product code" and rows[1][0] == "T1-0001" and len(rows) == 4
    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command("e2e_workbook", "retailers")


def test_e2e_ids_lists_seeded_records(settings, capsys):
    import json

    settings.DEBUG = True
    call_command("seed", "--no-photos")
    capsys.readouterr()
    call_command("e2e_ids")
    ids = json.loads(capsys.readouterr().out)
    assert set(ids) == {
        "tenant",
        "retailer",
        "product",
        "price_list",
        "rule",
        "import_job",
        "receipt",
        "draft_receipt",
        "adjustment",
        "shop_order",
        "order",
        "fulfilment",
        "backorder_product",
        "invoice",
        "credit_note",
        "payment",
        "refund",
        "shop_invoice",
        "returnable_shop_invoice",
        "ewaybill_invoice",
        "shop_checkout",
        "supplier",
        "purchase_order",
        "draft_purchase_order",
        "free_goods_scheme",
        "return_request",
    }
    assert ids["supplier"] and ids["purchase_order"] and ids["draft_purchase_order"]  # 9a
    assert ids["free_goods_scheme"]  # 9b
    assert ids["return_request"]  # 9c
    assert ids["returnable_shop_invoice"]  # 9c: a bill the shop can still return from
    assert ids["receipt"] and ids["draft_receipt"] and ids["adjustment"]  # from the demo stock
    assert ids["shop_order"] and ids["order"] and ids["fulfilment"] and ids["backorder_product"]
    assert ids["invoice"] and ids["credit_note"] and ids["payment"] and ids["refund"]  # billing
    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command("e2e_ids")


def test_seed_adds_demo_stock_once(settings):
    from apps.inventory.models import StockAdjustment, StockAlert, StockInward, StockLevel
    from apps.inventory.tests.helpers import check_invariants

    settings.DEBUG = True
    call_command("seed", "--no-photos")
    call_command("seed", "--no-photos")  # a second run adds nothing
    sharma = Tenant.objects.get(slug="sharma")
    with tenant_context(sharma.id):
        assert StockAdjustment.objects.count() == 3  # opening stock, shelf count, damage
        # Two demo receipts, the demo orders' receipt that serves a waiting backorder, and one
        # received against a demo purchase order (Phase 9a).
        assert StockInward.objects.filter(status="POSTED").count() == 4
        assert StockInward.objects.filter(purchase_order__isnull=False).count() == 1
        assert StockInward.objects.filter(status="DRAFT").count() == 1
        assert StockInward.objects.filter(cost_pending_lines__gt=0).count() == 2
        assert StockLevel.objects.filter(quantity_on_hand=0).exists()
        alert_types = set(
            StockAlert.objects.filter(status="OPEN").values_list("alert_type", flat=True)
        )
        assert alert_types == {"LOW_STOCK", "OUT_OF_STOCK", "BACKORDER_DEMAND"}
    check_invariants(sharma)


def test_seed_fills_every_order_tab_once(settings):
    from apps.orders.models import BackorderAllocation, Order
    from apps.orders.selectors import TABS
    from apps.orders.tests.helpers import check_order_invariants

    settings.DEBUG = True
    call_command("seed", "--no-photos")
    call_command("seed", "--no-photos")
    for slug in ("sharma", "patel"):
        tenant = Tenant.objects.get(slug=slug)
        with tenant_context(tenant.id):
            assert Order.objects.count() == 11, slug
            for tab, condition in TABS.items():
                assert Order.objects.filter(condition).exists(), (slug, tab)
            statuses = set(Order.objects.values_list("status", flat=True))
            assert statuses >= {"PLACED", "ON_HOLD", "ACCEPTED", "PACKED", "DISPATCHED"}
            assert {"COMPLETED", "REJECTED", "CANCELLED"} <= statuses
            assert Order.objects.filter(placed_via="STAFF", placed_by_label__endswith="(Sales)")
            assert BackorderAllocation.objects.filter(status="PROPOSED").count() == 1
            assert Order.objects.exclude(retailer_note="").exists()
        check_order_invariants(tenant)


def test_e2e_dispatch_sends_a_shipment_with_or_without_a_code(settings, capsys):
    """The delivery E2E's helper (ADR-057): a shipment on its way to Ganesh Kirana."""
    import json

    from apps.orders.models import Fulfilment

    settings.DEBUG = True
    call_command("seed", "--no-photos")
    capsys.readouterr()
    call_command("e2e_dispatch")
    plain = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    call_command("e2e_dispatch", "--with-code")
    coded = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert plain["code"] == "" and len(coded["code"]) == 4
    sharma = Tenant.objects.get(slug="sharma")
    with tenant_context(sharma.id):
        statuses = set(
            Fulfilment.objects.filter(pk__in=[plain["shipment"], coded["shipment"]]).values_list(
                "status", flat=True
            )
        )
        assert statuses == {"DISPATCHED"}
        from apps.platform.selectors import get_setting

        assert get_setting("orders.delivery_code", sharma.id) is False  # restored
