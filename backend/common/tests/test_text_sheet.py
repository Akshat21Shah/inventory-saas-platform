"""The translation sheet (ADR-060 item 11): every screen, Android app, server and notification
text with its Hindi and Marathi and their review status; a translator's copy is checked
(placeholders, plural forms, SMS and WhatsApp lengths) and written back into the files, with review
marks and a data migration for the platform's copies of changed notification texts. Works on
copies of the files."""

import json
import shutil
from collections.abc import Callable, Iterator
from io import BytesIO
from pathlib import Path
from typing import Any

import polib
import pytest
from openpyxl import load_workbook

from apps.notifications import catalog
from apps.notifications.catalog import texts_in
from apps.notifications.defaults import refresh_translated_templates
from apps.notifications.models import PlatformTemplate
from common import text_sheet
from common.text_sheet import NEEDS_REVIEW, REVIEWED, SHEET, icu_arguments
from common.translations import compiled

pytestmark = pytest.mark.django_db
BACKEND = Path(__file__).resolve().parents[2]
WELCOME_SMS = "notification:retailer.welcome/SHOP/SMS/body"


@pytest.fixture
def files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, settings: Any) -> Iterator[Path]:
    """Copies of the text files, so the tests write nothing real."""
    shutil.copytree(
        settings.WEB_MESSAGES_DIR, tmp_path / "web", ignore=shutil.ignore_patterns("*.ts")
    )
    settings.WEB_MESSAGES_DIR = tmp_path / "web"
    shutil.copytree(settings.APP_MESSAGES_DIR, tmp_path / "app")
    settings.APP_MESSAGES_DIR = tmp_path / "app"
    shutil.copytree(BACKEND / "locale", tmp_path / "locale")
    monkeypatch.setattr(
        text_sheet, "po_path", lambda code: tmp_path / "locale" / code / "LC_MESSAGES" / "django.po"
    )
    shutil.copytree(catalog.TEXTS_DIR, tmp_path / "catalog")
    monkeypatch.setattr(catalog, "TEXTS_DIR", tmp_path / "catalog")
    monkeypatch.setattr(text_sheet, "REVIEW_DIR", tmp_path / "reviewed")
    (tmp_path / "migrations").mkdir()
    for path in text_sheet.MIGRATIONS_DIR.glob("0*.py"):
        (tmp_path / "migrations" / path.name).write_text("")
    monkeypatch.setattr(text_sheet, "MIGRATIONS_DIR", tmp_path / "migrations")
    texts_in.cache_clear()
    yield tmp_path
    texts_in.cache_clear()


Edit = Callable[[Any, dict[str, int], dict[str, int]], None]


def _sheet(edit: Edit | None = None) -> BytesIO:
    """The sheet as exported, with ``edit(sheet, columns, rows)`` applied, as a file."""
    book = text_sheet.workbook()
    sheet = book[SHEET]
    columns = {str(c.value): int(c.column or 0) for c in sheet[1]}
    rows = {str(sheet.cell(r, 1).value): r for r in range(2, sheet.max_row + 1)}
    if edit:
        edit(sheet, columns, rows)
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def _set(key: str, column: str, value: str) -> Edit:
    def edit(sheet: Any, columns: dict[str, int], rows: dict[str, int]) -> None:
        sheet.cell(rows[key], columns[column]).value = value

    return edit


def _all(*edits: Edit) -> Edit:
    def edit(sheet: Any, columns: dict[str, int], rows: dict[str, int]) -> None:
        for one in edits:
            one(sheet, columns, rows)

    return edit


def test_the_sheet_has_every_text_with_its_status(files):
    book = load_workbook(_sheet())
    sheet = book[SHEET]
    header = [c.value for c in sheet[1]]
    assert header == [
        "Key", "Where", "English", "Hindi", "Hindi status", "Marathi", "Marathi status", "Note",
    ]  # fmt: skip
    keys = {str(sheet.cell(r, 1).value) for r in range(2, sheet.max_row + 1)}
    assert "screen:auth.signInTitle" in keys and "server:Enter your name." in keys
    assert WELCOME_SMS in keys
    assert "app:update.title" in keys  # the Android app's own texts (ADR-061)
    assert not any("/SUPPLIER/" in key for key in keys)  # suppliers' letters stay English
    assert sheet.max_row > 5000
    assert {sheet.cell(r, 5).value for r in range(2, 50)} == {NEEDS_REVIEW}
    assert "How to use" in book.sheetnames
    # Re-reading the sheet as it was made changes nothing.
    outcome = text_sheet.read(_sheet())
    assert (outcome.changes, outcome.errors, outcome.warnings) == ([], [], [])


def test_a_translators_copy_is_written_back_with_its_review_marks(files):
    new_sms = "{{ distributor }}: नमस्ते! ऑर्डर करें: {{ link }}"
    stream = _sheet(
        _all(
            _set("screen:auth.signInTitle", "Hindi", "लॉग इन करें"),
            _set("screen:auth.signInTitle", "Hindi status", REVIEWED),
            _set("server:Enter your name.", "Marathi", "कृपया तुमचे नाव लिहा."),
            _set("server:Enter your name.", "Marathi status", REVIEWED),
            _set(WELCOME_SMS, "Hindi", new_sms),
            _set(WELCOME_SMS, "Hindi status", REVIEWED),
            _set("app:update.button", "Marathi", "प्ले स्टोअर उघडा"),
        )
    )
    outcome = text_sheet.read(stream)
    assert outcome.errors == []
    assert {(c.key, c.code) for c in outcome.changes} == {
        ("screen:auth.signInTitle", "hi"),
        ("server:Enter your name.", "mr"),
        (WELCOME_SMS, "hi"),
        ("app:update.button", "mr"),
    }
    old_sms = texts_in("hi")["retailer.welcome"]["SHOP"]["SMS"].body
    text_sheet.apply(outcome)

    hindi = json.loads((files / "web" / "hi.json").read_text("utf-8"))
    assert hindi["auth"]["signInTitle"] == "लॉग इन करें"
    app = json.loads((files / "app" / "mr.json").read_text("utf-8"))
    assert app["update"]["button"] == "प्ले स्टोअर उघडा"
    marathi = polib.pofile(str(files / "locale" / "mr" / "LC_MESSAGES" / "django.po"))
    assert marathi.find("Enter your name.").msgstr == "कृपया तुमचे नाव लिहा."
    assert (files / "locale" / "mr" / "LC_MESSAGES" / "django.mo").read_bytes() == (
        compiled(marathi)
    )
    assert texts_in("hi")["retailer.welcome"]["SHOP"]["SMS"].body == new_sms
    marks = json.loads((files / "reviewed" / "hi.json").read_text("utf-8"))
    assert set(marks) == {"screen:auth.signInTitle", WELCOME_SMS}
    migration = next((files / "migrations").glob("*_reviewed_texts_*.py")).read_text("utf-8")
    assert "refresh_translated_templates" in migration and repr(old_sms) in migration
    # Now shown as reviewed, until its English changes.
    progress = {p.code: p for p in text_sheet.progress()}
    assert (progress["hi"].reviewed, progress["mr"].reviewed) == (2, 1)
    assert progress["hi"].translated == progress["hi"].total
    sheet = load_workbook(_sheet())[SHEET]
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == WELCOME_SMS)
    assert (sheet.cell(row, 4).value, sheet.cell(row, 5).value) == (new_sms, REVIEWED)


@pytest.mark.parametrize(
    ("key", "column", "text", "problem"),
    [
        ("screen:nav.itemCount", "Hindi", "कार्ट में सामान", "Keep as they are: {count}"),
        (
            "screen:nav.itemCount",
            "Hindi",
            "{count, plural, one {# सामान}} कार्ट में",
            "with an “other” form",
        ),
        ("server:Row %(row)s, column “%(column)s”: %(message)s", "Marathi", "ओळ", "%(row)s"),
        (WELCOME_SMS, "Hindi", "{{ distributor }}: " + "क" * 80 + " {{ link }}", "one part"),
        (WELCOME_SMS, "Hindi", "स्वागत है! {{ link }}", "start with"),
        (WELCOME_SMS, "Marathi", "{{ distributor }}: {{ secret }}", "use only these values"),
        ("app:account.version", "Marathi", "ॲप आवृत्ती {version}", "{build}"),
        ("screen:auth.signInTitle", "Hindi status", REVIEWED, None),  # fine
    ],
)
def test_mistakes_are_refused(files, key, column, text, problem):
    outcome = text_sheet.read(_sheet(_set(key, column, text)))
    if problem is None:
        assert outcome.errors == []
        return
    assert any(problem in error for error in outcome.errors), outcome.errors
    with pytest.raises(ValueError):
        text_sheet.apply(outcome)


def test_stale_and_unknown_rows_are_skipped_and_empty_reviewed_texts_refused(files):
    outcome = text_sheet.read(
        _sheet(
            _all(
                _set("screen:auth.signInTitle", "English", "Log in"),  # the English changed
                _set("screen:auth.signInTitle", "Hindi", "लॉग इन"),
                _set("server:Enter your name.", "Key", "server:Gone."),
                _set(WELCOME_SMS, "Hindi", ""),
                _set(WELCOME_SMS, "Hindi status", REVIEWED),
            )
        )
    )
    assert outcome.changes == []
    assert any("English changed" in w for w in outcome.warnings)
    assert any("no longer in the app" in w for w in outcome.warnings)
    assert any("marked reviewed but empty" in e for e in outcome.errors)


def test_placeholders_are_read_from_the_message_format():
    assert icu_arguments("{count, plural, =0 {Products} one {# product} other {# products}}") == [
        "count"
    ]
    assert icu_arguments("Hello {name}, {n, number} left; {when, date, short}") == [
        "name",
        "n",
        "when",
    ]
    assert icu_arguments("{gender, select, male {He has {count}} other {They have {count}}}") == [
        "gender",
        "count",
        "count",
    ]
    assert icu_arguments("No placeholders") == []


def test_the_platforms_copies_take_reviewed_words_unless_the_super_admin_changed_them(files):
    welcome = PlatformTemplate.objects.get(
        event_code="retailer.welcome", audience="SHOP", channel="SMS", locale="hi"
    )
    whatsapp = PlatformTemplate.objects.filter(channel="WHATSAPP", locale="hi").first()
    assert whatsapp is not None
    whatsapp.approval_status = "APPROVED"
    whatsapp.save()
    edited = PlatformTemplate.objects.get(
        event_code="retailer.welcome", audience="SHOP", channel="EMAIL", locale="hi"
    )
    edited.body = "सुपर एडमिन के शब्द"
    edited.save()
    old = {
        row.pk: (row.subject, row.body) for row in (welcome, whatsapp)
    }  # what the catalogue said
    data = json.loads((files / "catalog" / "hi.json").read_text("utf-8"))
    data["retailer.welcome"]["SHOP"]["SMS"]["body"] = "{{ distributor }}: नया स्वागत {{ link }}"
    data["retailer.welcome"]["SHOP"]["EMAIL"]["body"] = "नया ईमेल {{ shop }} {{ link }}"
    wa = data[whatsapp.event_code][whatsapp.audience]["WHATSAPP"]
    wa["body"] = wa["body"] + " ।"
    (files / "catalog" / "hi.json").write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    texts_in.cache_clear()
    changed = [
        ("retailer.welcome", "SHOP", "SMS", "hi", *old[welcome.pk]),
        ("retailer.welcome", "SHOP", "EMAIL", "hi", edited.subject, "the catalogue's old words"),
        (whatsapp.event_code, whatsapp.audience, "WHATSAPP", "hi", *old[whatsapp.pk]),
    ]
    assert refresh_translated_templates(PlatformTemplate, changed) == 2
    welcome.refresh_from_db()
    edited.refresh_from_db()
    whatsapp.refresh_from_db()
    assert welcome.body == "{{ distributor }}: नया स्वागत {{ link }}"
    assert edited.body == "सुपर एडमिन के शब्द"  # the super admin's own words stay
    assert whatsapp.body.endswith(" ।") and whatsapp.approval_status == "NOT_SUBMITTED"
