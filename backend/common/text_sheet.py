"""Every text the app shows, in one spreadsheet for translators (ADR-060 item 11).

``manage.py texts_export`` writes it, and the super admin downloads the same sheet. A native
speaker corrects the Hindi or Marathi and marks each text "reviewed"; ``manage.py texts_import``
checks their copy (placeholders, plural forms, SMS and WhatsApp lengths) and, with ``--apply``,
writes it back into the files, which are then committed:

- screens: ``web/messages/<code>.json`` (keys ``screen:<dotted key>``);
- the Android app's own texts: ``mobile/messages/app/<code>.json`` (keys ``app:<dotted key>``;
  the app shows the screens' texts too, ADR-061);
- server messages: ``locale/<code>/LC_MESSAGES/django.po`` and its compiled ``.mo``
  (``server:<English>``, ``server:<context>|<English>``; a plural's forms ``…#0``, ``…#1``);
- notification texts: ``apps/notifications/catalog_texts/<code>.json``
  (``notification:<event>/<audience>/<channel>/<subject|body>``); a data migration written with
  them gives the platform's copies the super admin hasn't changed the new words.

A text counts as reviewed once a sheet marks it so, until its English changes
(``locale/reviewed/<code>.json`` keeps a fingerprint of the English each was reviewed against).
Column headings and notes stay English (owner)."""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import IO, Any

import polib
from django.conf import settings
from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import languages
from common.translations import BRACES, PERCENT, compiled, placeholders, po_path, po_text

REVIEWED, NEEDS_REVIEW = "reviewed", "needs review"
SHEET = "Texts"
BACKEND = Path(__file__).resolve().parents[1]
REVIEW_DIR = BACKEND / "locale" / "reviewed"
MIGRATIONS_DIR = BACKEND / "apps" / "notifications" / "migrations"
WHATSAPP_LIMIT = 1024  # characters in a WhatsApp template's body
CHANNELS = {"IN_APP": "in the app", "EMAIL": "email", "SMS": "SMS", "WHATSAPP": "WhatsApp"}
SAMPLE_DISTRIBUTOR = "Sharma Traders"

ICU_FORMS = ("plural", "select", "selectordinal")
ICU_PLURAL = re.compile(r",\s*(plural|select|selectordinal)\s*,")
TAG = re.compile(r"<(/?)([A-Za-z]\w*)\s*/?>")
VARIABLE = re.compile(r"{{\s*(\w+)\s*}}")


@dataclass
class Row:
    key: str
    where: str
    english: str
    texts: dict[str, str]  # by language code, English left out
    note: str = ""

    @property
    def kind(self) -> str:
        return self.key.split(":", 1)[0]


def translated() -> list[languages.Language]:
    """Every language but English."""
    return [lang for lang in languages.all_languages() if lang.code != languages.DEFAULT]


def fingerprint(english: str) -> str:
    return hashlib.sha256(english.encode("utf-8")).hexdigest()[:12]


def review_marks(code: str) -> dict[str, str]:
    """Each reviewed text's key and the fingerprint of the English it was reviewed against."""
    path = REVIEW_DIR / f"{code}.json"
    marks: dict[str, str] = json.loads(path.read_text("utf-8")) if path.exists() else {}
    return marks


def status(row: Row, code: str, marks: dict[str, str]) -> str:
    return REVIEWED if marks.get(row.key) == fingerprint(row.english) else NEEDS_REVIEW


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text("utf-8"))


def _write_json(path: Path, data: Any, *, sort_keys: bool = False) -> None:
    text = json.dumps(data, ensure_ascii=False, indent=2, sort_keys=sort_keys)
    path.write_text(text + "\n", "utf-8")


# --- The texts --------------------------------------------------------------------------------


def _flat(tree: dict[str, Any], prefix: str = "") -> dict[str, str]:
    out: dict[str, str] = {}
    for key, value in tree.items():
        if isinstance(value, dict):
            out.update(_flat(value, f"{prefix}{key}."))
        else:
            out[f"{prefix}{key}"] = str(value)
    return out


# Texts kept as JSON message files, by key prefix: where they are and what the sheet calls them.
JSON_TEXTS = {"screen": "Screen", "app": "Android app"}


def message_file(kind: str, code: str) -> Path:
    folder = settings.WEB_MESSAGES_DIR if kind == "screen" else settings.APP_MESSAGES_DIR
    return Path(folder) / f"{code}.json"


def _server_tokens(english: str) -> list[str]:
    """A server message's placeholders as written: %(name)s, {name}."""
    return [m.group(0) for m in PERCENT.finditer(english)] + [
        m.group(0) for m in BRACES.finditer(english)
    ]


def _keep(found: Iterable[str]) -> str:
    listed = list(dict.fromkeys(found))
    return f"Keep as they are: {', '.join(listed)}." if listed else ""


def _screen_note(english: str) -> str:
    parts = [_keep(f"{{{a}}}" for a in icu_arguments(english))]
    if TAG.search(english):
        parts.append("Keep the <tags> around the same words.")
    if ICU_PLURAL.search(english):
        parts.append("Give every form (one, other); only translate the words inside { }.")
    return " ".join(p for p in parts if p)


def _json_rows(kind: str, codes: list[str]) -> list[Row]:
    if not message_file(kind, "en").exists():
        return []
    english = _flat(_read_json(message_file(kind, "en")))
    own = {
        code: _flat(_read_json(message_file(kind, code)))
        if message_file(kind, code).exists()
        else {}
        for code in codes
    }
    return [
        Row(
            f"{kind}:{key}",
            JSON_TEXTS[kind],
            text,
            {c: own[c].get(key, "") for c in codes},
            _screen_note(text),
        )
        for key, text in english.items()
    ]


def _catalog(code: str) -> polib.POFile | None:
    path = po_path(code)
    return polib.pofile(str(path)) if path.exists() else None


def _server_key(entry: polib.POEntry) -> str:
    return f"server:{entry.msgctxt}|{entry.msgid}" if entry.msgctxt else f"server:{entry.msgid}"


def _server_rows(codes: list[str]) -> list[Row]:
    rows: dict[str, Row] = {}
    for code in codes:
        catalog = _catalog(code)
        if catalog is None:
            continue
        for entry in catalog:
            if entry.obsolete:
                continue
            base = _server_key(entry)
            forms = (
                [(f"{base}#{i}", entry.msgid if i == 0 else entry.msgid_plural, text)
                 for i, text in sorted(entry.msgstr_plural.items())]
                if entry.msgid_plural
                else [(base, entry.msgid, entry.msgstr)]
            )  # fmt: skip
            for key, english, text in forms:
                row = rows.get(key)
                if row is None:
                    note = _keep(placeholders(english))
                    row = rows[key] = Row(
                        key, "Server message", english, dict.fromkeys(codes, ""), note
                    )
                row.texts[code] = text
    return list(rows.values())


def _notification_note(channel: str, english: str) -> str:
    parts = [_keep(dict.fromkeys(f"{{{{ {v} }}}}" for v in VARIABLE.findall(english)))]
    if channel == "SMS":
        parts.append(
            "SMS: Hindi and Marathi take 70 characters a part (each part is paid for); keep it "
            "short. The welcome SMS must fit one part."
        )
    if channel == "WHATSAPP":
        parts.append(
            f"WhatsApp: at most {WHATSAPP_LIMIT:,} characters; a changed text needs WhatsApp's "
            "approval again."
        )
    if channel in ("SMS", "WHATSAPP"):
        parts.append("Start with {{ distributor }}: as the English does.")
    return " ".join(p for p in parts if p)


def _notification_rows(codes: list[str]) -> list[Row]:
    from apps.notifications.catalog import DEFAULT_TEXTS, texts_in
    from apps.notifications.models import Audience

    rows = []
    for event, audiences in DEFAULT_TEXTS.items():
        for audience, channels in audiences.items():
            if audience == Audience.SUPPLIER:  # suppliers' letters stay English (owner)
                continue
            for channel, text in channels.items():
                for part in ("subject", "body"):
                    english = getattr(text, part)
                    if not english:
                        continue
                    own = {}
                    for code in codes:
                        found = texts_in(code).get(event, {}).get(audience, {}).get(channel)
                        own[code] = getattr(found, part) if found else ""
                    rows.append(
                        Row(
                            f"notification:{event}/{audience}/{channel}/{part}",
                            f"Notification ({CHANNELS.get(channel, channel)})",
                            english,
                            own,
                            _notification_note(channel, english),
                        )
                    )
    return rows


def rows() -> list[Row]:
    """Every text: screens, the Android app's own, server messages, notifications."""
    codes = [lang.code for lang in translated()]
    screens = _json_rows("screen", codes) + _json_rows("app", codes)
    return screens + _server_rows(codes) + _notification_rows(codes)


@dataclass(frozen=True)
class Progress:
    code: str
    total: int
    translated: int
    reviewed: int


def progress() -> list[Progress]:
    """How many texts each language has, and how many a native speaker has reviewed."""
    every = rows()
    found = []
    for lang in translated():
        marks = review_marks(lang.code)
        found.append(
            Progress(
                lang.code,
                len(every),
                sum(1 for r in every if r.texts.get(lang.code)),
                sum(1 for r in every if status(r, lang.code, marks) == REVIEWED),
            )
        )
    return found


# --- The spreadsheet --------------------------------------------------------------------------


def _text(cell: Cell, value: str) -> None:
    """A cell that holds text, even text starting with "=" (never a formula)."""
    cell.value = value
    cell.data_type = "s"


HOW_TO = (
    "How to review these texts",
    "",
    "1. Each row is one text the app shows. Read the English, then correct the Hindi or Marathi "
    "next to it.",
    "2. When a text is right, set its status to “reviewed”. Leave “needs "
    "review” for anything you are unsure of.",
    "3. Keep everything shown in the Note column exactly as it is: {name}, %(name)s, "
    "{{ name }}, <b> and </b>. They are filled in by the app.",
    "4. Don't change the Key, Where or English columns, and don't add or remove rows.",
    "5. Names of products, shops and people are never translated; amounts use the digits 0-9.",
    "6. Send the sheet back as it is (.xlsx).",
)


def workbook(found: list[Row] | None = None) -> Workbook:
    found = rows() if found is None else found
    langs = translated()
    marks = {lang.code: review_marks(lang.code) for lang in langs}
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = SHEET
    header = ["Key", "Where", "English"]
    for lang in langs:
        header += [lang.name, f"{lang.name} status"]
    header.append("Note")
    sheet.append(header)
    for row in found:
        values = [row.key, row.where, row.english]
        for lang in langs:
            values += [row.texts.get(lang.code, ""), status(row, lang.code, marks[lang.code])]
        values.append(row.note)
        sheet.append(values)
    wrap = Alignment(wrap_text=True, vertical="top")
    for column in sheet.iter_cols(min_row=2, max_row=sheet.max_row):
        for cell in column:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                _text(cell, cell.value)
            cell.alignment = wrap
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "D2"
    widths = [28, 16, 48] + [48, 14] * len(langs) + [40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    choice = DataValidation(type="list", formula1=f'"{NEEDS_REVIEW},{REVIEWED}"', allow_blank=False)
    sheet.add_data_validation(choice)
    for index in range(len(langs)):
        letter = get_column_letter(5 + 2 * index)
        choice.add(f"{letter}2:{letter}{sheet.max_row}")
    guide = book.create_sheet("How to use")
    for line in HOW_TO:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True)
    return book


def export(stream: IO[bytes]) -> int:
    """Write the sheet; returns how many texts it has."""
    found = rows()
    workbook(found).save(stream)
    return len(found)


# --- Reading a translator's copy --------------------------------------------------------------


@dataclass(frozen=True)
class Change:
    key: str
    code: str
    before: str
    after: str


@dataclass
class Outcome:
    changes: list[Change] = field(default_factory=list)
    reviewed: dict[str, dict[str, str]] = field(default_factory=lambda: defaultdict(dict))
    unreviewed: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    rows_read: int = 0


def _sample(text: str) -> str:
    from apps.notifications.render import substitute
    from apps.notifications.texts import SAMPLE

    return substitute(text, {**SAMPLE, "distributor": SAMPLE_DISTRIBUTOR})


def icu_arguments(text: str) -> list[str]:
    """The arguments of an ICU message, nested ones too, but not the words of a plural's or a
    select's forms: "{count, plural, =0 {Products} other {# products}}" → ["count"]."""
    found: list[str] = []

    def message(i: int, nested: bool) -> int:
        while i < len(text):
            if text[i] == "{":
                i = argument(i + 1)
            elif text[i] == "}" and nested:
                return i + 1
            else:
                i += 1
        return i

    def argument(i: int) -> int:
        end = i
        while end < len(text) and text[end] not in ",}":
            end += 1
        found.append(text[i:end].strip())
        if end >= len(text) or text[end] == "}":
            return end + 1
        kind_end = end + 1
        while kind_end < len(text) and text[kind_end] not in ",}":
            kind_end += 1
        kind = text[end + 1 : kind_end].strip()
        if kind_end < len(text) and text[kind_end] == "," and kind in ICU_FORMS:
            i = kind_end + 1
            while i < len(text):  # each form: a selector, then {its message}
                while i < len(text) and text[i] not in "{}":
                    i += 1
                if i >= len(text) or text[i] == "}":
                    return i + 1
                i = message(i + 1, nested=True)
            return i
        while kind_end < len(text) and text[kind_end] != "}":  # {n, number}, {d, date, …}
            kind_end += 1
        return kind_end + 1

    message(0, nested=False)
    return found


def _check_screen(english: str, text: str) -> list[str]:
    found = []
    expected = icu_arguments(english)
    if sorted(set(icu_arguments(text))) != sorted(set(expected)):
        found.append(_keep(f"{{{a}}}" for a in expected) or "use no {placeholders}")
    if sorted(TAG.findall(text)) != sorted(TAG.findall(english)):
        found.append("keep the same <tags>")
    for kind in ICU_PLURAL.findall(english):
        if not re.search(rf",\s*{kind}\s*,", text) or not re.search(r"\bother\s*\{", text):
            found.append(f"keep the {kind} forms, with an “other” form")
    if text.count("{") != text.count("}"):
        found.append("a { or } is missing")
    return found


def _check_notification(row: Row, text: str, before: str) -> tuple[list[str], list[str]]:
    from apps.notifications import sms_length

    event, _audience, channel, part = row.key.split(":", 1)[1].split("/")
    errors, warnings = [], []
    used, expected = set(VARIABLE.findall(text)), set(VARIABLE.findall(row.english))
    if channel == "SMS" and part == "body":
        if not used <= expected:  # short enough for one part: may leave some out
            errors.append(f"use only these values: {', '.join(sorted(expected))}")
    elif used != expected:
        errors.append(f"keep the values: {', '.join(sorted(expected))}")
    if (
        channel in ("SMS", "WHATSAPP")
        and part == "body"
        and not text.startswith("{{ distributor }}: ")
    ):
        errors.append("start with “{{ distributor }}: ” as the English does")
    if channel == "WHATSAPP" and len(text) > WHATSAPP_LIMIT:
        errors.append(f"{len(text):,} characters; WhatsApp allows {WHATSAPP_LIMIT:,}")
    if channel == "WHATSAPP" and text != before:
        warnings.append("a changed WhatsApp text needs WhatsApp's approval again")
    if channel == "SMS" and part == "body":
        parts = sms_length.parts(_sample(text))
        if event == "retailer.welcome" and parts > 1:
            length = sms_length.length(_sample(text))
            errors.append(
                f"the welcome SMS must fit one part; it takes {parts} ({length} characters)"
            )
        elif before and parts > sms_length.parts(_sample(before)):
            warnings.append(f"now takes {parts} SMS parts, was {sms_length.parts(_sample(before))}")
    return errors, warnings


def check_text(row: Row, text: str, before: str) -> tuple[list[str], list[str]]:
    """What is wrong with ``text`` as ``row``'s translation (errors, warnings)."""
    if row.kind in JSON_TEXTS:
        return _check_screen(row.english, text), []
    if row.kind == "server":
        if placeholders(text) != placeholders(row.english):
            return [_keep(_server_tokens(row.english)) or "use no placeholders"], []
        return [], []
    return _check_notification(row, text, before)


def _cell(value: Any) -> str:
    return "" if value is None else str(value).replace("\r\n", "\n")


def read(stream: IO[bytes] | Path) -> Outcome:
    """Check a translator's sheet against today's texts; nothing is written."""
    outcome = Outcome()
    book = load_workbook(stream, read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        outcome.errors.append(f"the sheet has no “{SHEET}” tab")
        return outcome
    lines = book[SHEET].iter_rows(values_only=True)
    header = [_cell(v).strip() for v in next(lines, ())]
    column = {name: index for index, name in enumerate(header)}
    langs = [lang for lang in translated() if lang.name in column]
    if "Key" not in column or "English" not in column or not langs:
        outcome.errors.append("the first row must be the headings this app wrote")
        return outcome
    current = {row.key: row for row in rows()}
    for number, values in enumerate(lines, start=2):
        key = _cell(values[column["Key"]]) if len(values) > column["Key"] else ""
        if not key.strip():  # kept exactly: a server message's key is its English
            continue
        outcome.rows_read += 1
        row = current.get(key)
        if row is None:
            outcome.warnings.append(f"row {number}: {key} is no longer in the app; skipped")
            continue
        if _cell(values[column["English"]]) != row.english:
            outcome.warnings.append(
                f"row {number}: the English changed since this sheet was made; skipped (export "
                "a new sheet for it)"
            )
            continue
        for lang in langs:
            text = _cell(values[column[lang.name]])
            state = _cell(values[column.get(f"{lang.name} status", -1)]).strip().lower()
            where = f"row {number}, {lang.name}"
            if not text.strip():
                if state == REVIEWED:
                    outcome.errors.append(f"{where}: marked reviewed but empty")
                continue
            before = row.texts.get(lang.code, "")
            errors, warnings = check_text(row, text, before)
            outcome.errors += [f"{where}: {e}" for e in errors]
            outcome.warnings += [f"{where}: {w}" for w in warnings]
            if errors:
                continue
            if text != before:
                outcome.changes.append(Change(key, lang.code, before, text))
            if state == REVIEWED:
                outcome.reviewed[lang.code][key] = fingerprint(row.english)
            elif state == NEEDS_REVIEW:
                outcome.unreviewed[lang.code].add(key)
            elif state:
                outcome.warnings.append(f"{where}: status “{state}” isn't one of the two")
    return outcome


# --- Writing it back --------------------------------------------------------------------------


def _set_path(tree: dict[str, Any], dotted: str, value: str) -> None:
    parts = dotted.split(".")
    for part in parts[:-1]:
        tree = tree.setdefault(part, {})
    tree[parts[-1]] = value


def _apply_json(kind: str, changes: list[Change]) -> list[Path]:
    written = []
    for code in sorted({c.code for c in changes}):
        path = message_file(kind, code)
        data = _read_json(path)
        for change in (c for c in changes if c.code == code):
            _set_path(data, change.key.split(":", 1)[1], change.after)
        _write_json(path, data)
        written.append(path)
    return written


def _apply_server(changes: list[Change]) -> list[Path]:
    written = []
    for code in sorted({c.code for c in changes}):
        catalog = _catalog(code)
        assert catalog is not None
        entries = {_server_key(e): e for e in catalog if not e.obsolete}
        for change in (c for c in changes if c.code == code):
            base, _, form = change.key.partition("#")
            entry = entries[base]
            if form:
                entry.msgstr_plural[int(form)] = change.after
            else:
                entry.msgstr = change.after
        path = po_path(code)
        path.write_text(po_text(catalog), encoding="utf-8")
        path.with_suffix(".mo").write_bytes(compiled(catalog))
        written += [path, path.with_suffix(".mo")]
    return written


def _apply_notifications(changes: list[Change]) -> list[Path]:
    from apps.notifications.catalog import TEXTS_DIR, texts_in

    written = []
    refreshed: list[tuple[str, str, str, str, str, str]] = []
    for code in sorted({c.code for c in changes}):
        path = TEXTS_DIR / f"{code}.json"
        data = _read_json(path)
        mine = [c for c in changes if c.code == code]
        for event, audience, channel in sorted(
            {tuple(c.key.split(":", 1)[1].split("/")[:3]) for c in mine}
        ):
            text = data.setdefault(event, {}).setdefault(audience, {}).setdefault(channel, {})
            old_subject, old_body = text.get("subject", ""), text.get("body", "")
            refreshed.append((event, audience, channel, code, old_subject, old_body))
            for change in mine:
                if change.key.split(":", 1)[1].startswith(f"{event}/{audience}/{channel}/"):
                    text[change.key.rsplit("/", 1)[1]] = change.after
            text.setdefault("subject", "")
        _write_json(path, data, sort_keys=True)
        written.append(path)
    texts_in.cache_clear()
    if refreshed:
        written.append(_refresh_migration(refreshed))
    return written


def _refresh_migration(changed: list[tuple[str, str, str, str, str, str]]) -> Path:
    """A data migration that gives the platform's copies of these texts the new words, where
    they still hold the old ones (the super admin's own edits are kept)."""
    names = sorted(p.stem for p in MIGRATIONS_DIR.glob("[0-9][0-9][0-9][0-9]_*.py"))
    last = names[-1]
    path = MIGRATIONS_DIR / f"{int(last[:4]) + 1:04d}_reviewed_texts_{date.today():%Y%m%d}.py"
    listed = "\n".join(f"    {item!r}," for item in changed)
    path.write_text(
        f'''"""Translations a native speaker reviewed (texts_import, ADR-060 item 11): the
platform's copies that still hold the old words take the catalogue's new ones; the super admin's
own edits are kept, and a changed WhatsApp text goes back to "not submitted"."""

from django.db import migrations

from apps.notifications.defaults import refresh_translated_templates

# event, audience, channel, language, the old subject, the old body
CHANGED = [
{listed}
]


def forwards(apps, schema_editor):
    refresh_translated_templates(apps.get_model("notifications", "PlatformTemplate"), CHANGED)


class Migration(migrations.Migration):
    dependencies = [("notifications", "{last}")]

    operations = [migrations.RunPython(forwards, migrations.RunPython.noop)]
''',
        encoding="utf-8",
    )
    return path


def _apply_marks(outcome: Outcome) -> list[Path]:
    written = []
    for code in sorted(set(outcome.reviewed) | set(outcome.unreviewed)):
        marks = review_marks(code)
        marks.update(outcome.reviewed.get(code, {}))
        for key in outcome.unreviewed.get(code, set()):
            marks.pop(key, None)
        REVIEW_DIR.mkdir(parents=True, exist_ok=True)
        path = REVIEW_DIR / f"{code}.json"
        _write_json(path, dict(sorted(marks.items())))
        written.append(path)
    return written


def apply(outcome: Outcome) -> list[Path]:
    """Write a checked sheet's changes and review marks into the files (dev, then commit)."""
    if outcome.errors:
        raise ValueError("the sheet has errors; nothing was written")
    by_kind: dict[str, list[Change]] = defaultdict(list)
    for change in outcome.changes:
        by_kind[change.key.split(":", 1)[0]].append(change)
    written = _apply_json("screen", by_kind["screen"]) + _apply_json("app", by_kind["app"])
    written += _apply_server(by_kind["server"])
    written += _apply_notifications(by_kind["notification"])
    written += _apply_marks(outcome)
    return written
